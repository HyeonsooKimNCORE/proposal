"""나어릴때 제안서 v3 — v2 보충 (작업지시서_v2보충.md 2단계)

실행: python3 build_v3.py  →  output/나어릴때_제안서_v3.pptx
- 기준: input/나어릴때_제안서_v2.pptx (원본은 읽기만 함)
- 새 장표는 v2의 기준 장표를 복제한 뒤 텍스트·표 내용과 배치만 바꿈
- 추가 원고: 나어릴때_제안서_추가원고.md / 표현 수정 11건(사용자 승인: 전부)
"""
import copy
import re
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).parent
SRC = ROOT / "input" / "나어릴때_제안서_v2.pptx"
OUT = ROOT / "output" / "나어릴때_제안서_v3.pptx"
ASSETS = ROOT / "assets"

# v1/v2와 같은 색·폰트
TEXT, SUB, CAP, LINE, WHITE = "211D1E", "5E5759", "8F8789", "BFBFBF", "FFFFFF"
LILAC, LILAC_M, LILAC_T = "7B2FB5", "C9A3E3", "F1E4F8"
PEACH, PEACH_T, IMG_BG, CHECK = "F6D2A8", "FEF2DA", "E6E1DE", "FFF176"
F_BODY, F_MED, F_BLACK = "Pretendard", "Pretendard Medium", "Pretendard Black"
X0, W = 0.45, 12.43

prs = Presentation(str(SRC))
SLD = prs.slides._sldIdLst
LOG = {"edits": [], "missing": set(), "checks": []}


def rgb(h):
    return RGBColor.from_string(h)


# ---------------------------------------------------------------- 텍스트
CHECK_MARK = "〔확인〕"


def split_check(text):
    out, i = [], 0
    while True:
        j = text.find(CHECK_MARK, i)
        if j < 0:
            out.append((text[i:], False))
            break
        out.append((text[i:j], False))
        k = j + len(CHECK_MARK)
        if k < len(text) and text[k].isdigit():
            m = re.search(r"\s", text[k:])
            e = k + m.start() if m else len(text)
        else:
            cands = [p for p in (text.find(" / ", k), text.find(")", k)) if p >= 0]
            e = min(cands) if cands else len(text)
        out.append((text[j:e], True))
        i = e
    return [(t, c) for t, c in out if t]


def set_run_font(run, size, bold=False, color=TEXT, font=F_BODY, highlight=False):
    if highlight:
        color = TEXT  # 노란 하이라이트 위 글자는 어두운 색으로 (흰 글자 대비 문제)
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = rgb(color)
    f.name = font
    rPr = run._r.get_or_add_rPr()
    latin = rPr.find(qn("a:latin"))
    ea = etree.SubElement(rPr, qn("a:ea"))
    latin.addnext(ea)
    ea.set("typeface", font)
    cs = etree.SubElement(rPr, qn("a:cs"))
    ea.addnext(cs)
    cs.set("typeface", font)
    if highlight:
        hl = etree.Element(qn("a:highlight"))
        etree.SubElement(hl, qn("a:srgbClr")).set("val", CHECK)
        rPr.find(qn("a:solidFill")).addnext(hl)
        LOG["checks"].append(run.text)


def fill_tf(tf, paras, size=14, color=TEXT, font=F_BODY, bold=False, align=PP_ALIGN.LEFT,
            spacing=1.15, space_after=0, bullet=False):
    if isinstance(paras, str):
        paras = [paras]
    tf.clear()
    for pi, para in enumerate(paras):
        p = tf.paragraphs[0] if pi == 0 else tf.add_paragraph()
        for r in list(p.runs):
            p._p.remove(r._r)
        p.alignment = align
        p.line_spacing = spacing
        if space_after:
            p.space_after = Pt(space_after)
        pPr = p._p.get_or_add_pPr()
        pPr.set("latinLnBrk", "0")
        pPr.set("eaLnBrk", "1")
        for old in pPr.findall(qn("a:buChar")) + pPr.findall(qn("a:buNone")):
            pPr.remove(old)
        runs = [(para, {})] if isinstance(para, str) else para
        for text, opts in runs:
            for seg, chk in split_check(text):
                r = p.add_run()
                r.text = seg
                set_run_font(r, opts.get("size", size), opts.get("bold", bold), opts.get("color", color),
                             opts.get("font", font), highlight=chk)
        if bullet:
            sz = runs[0][1].get("size", size)
            pPr.set("marL", str(int(Pt(sz))))
            pPr.set("indent", str(-int(Pt(sz))))
            etree.SubElement(pPr, qn("a:buChar")).set("char", "•")
        else:
            pPr.attrib.pop("marL", None)
            pPr.attrib.pop("indent", None)
            etree.SubElement(pPr, qn("a:buNone"))


def text_width(text, size):
    w = 0
    for ch in text:
        if "가" <= ch <= "힣" or ch in "〔〕":
            w += 0.95
        elif ch == " ":
            w += 0.28
        else:
            w += 0.58
    return w * size / 72


def fit_size(text, width=W, start=24, minimum=20):
    s = start
    while s > minimum and text_width(text, s) > width - 0.1:
        s -= 1
    return s


# ---------------------------------------------------------------- 도형
def _strip_style(shape):
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


def tb(slide, x, y, w, h, paras, size=14, anchor=MSO_ANCHOR.TOP, **kw):
    b = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = b.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", 0)
    tf.vertical_anchor = anchor
    fill_tf(tf, paras, size=size, **kw)
    return b


def box(slide, x, y, w, h, fill=WHITE, line=None, radius=0.08, shape=MSO_SHAPE.ROUNDED_RECTANGLE, line_w=1.0, name=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    _strip_style(s)
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(line_w)
    else:
        s.line.fill.background()
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = min(radius / min(w, h), 0.5)
    if name:
        s.name = name
    return s


def shape_text(s, paras, size=14, anchor=MSO_ANCHOR.MIDDLE, margin=0.12, **kw):
    tf = s.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    tf.vertical_anchor = anchor
    fill_tf(tf, paras, size=size, **kw)


def chip(slide, x, y, w, text, fill=LILAC_T, color=LILAC, size=11, h=0.3):
    s = box(slide, x, y, w, h, fill=fill, radius=h / 2)
    shape_text(s, [[(text, {"bold": True})]], size=size, color=color, align=PP_ALIGN.CENTER, margin=0.04)
    return s


def image(slide, x, y, w, h, fname):
    p = ASSETS / fname
    if p.exists():
        from PIL import Image as PImage
        iw, ih = PImage.open(p).size
        r = min(w / iw, h / ih)
        pw, ph = iw * r, ih * r
        pic = slide.shapes.add_picture(str(p), Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2), Inches(pw), Inches(ph))
        pic.name = f"IMG_{fname}"
        return pic
    LOG["missing"].add(fname)
    s = box(slide, x, y, w, h, fill=IMG_BG, radius=0.05, name=f"IMG_{fname}")
    shape_text(s, [[(f"[이미지 필요: {fname}]", {})]], size=12, color=SUB, align=PP_ALIGN.CENTER)
    return s


def place(shape, x=None, y=None, w=None, h=None):
    if x is not None:
        shape.left = Inches(x)
    if y is not None:
        shape.top = Inches(y)
    if w is not None:
        shape.width = Inches(w)
    if h is not None:
        shape.height = Inches(h)


def _next_id(slide):
    return max(int(e.get("id")) for e in slide.shapes._spTree.iter(qn("p:cNvPr"))) + 1


def clone(slide, shape):
    """같은 장 안에서 도형 복제 (id 중복 방지)"""
    el = copy.deepcopy(shape._element)
    el.find(".//" + qn("p:cNvPr")).set("id", str(_next_id(slide)))
    slide.shapes._spTree.append(el)
    return slide.shapes[-1]


def remove(shape):
    shape._element.getparent().remove(shape._element)


def notes(slide, text):
    if text and text != "-":
        slide.notes_slide.notes_text_frame.text = text


# ---------------------------------------------------------------- 장표 복제·이동
def find_slide(headline):
    for s in prs.slides:
        for sh in s.shapes:
            if sh.has_text_frame and headline in sh.text_frame.text:
                return s
    raise KeyError(headline)


def duplicate(src):
    new = prs.slides.add_slide(src.slide_layout)
    tree = new.shapes._spTree
    for el in list(tree):
        if el.tag not in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")):
            tree.remove(el)
    for el in src.shapes._spTree:
        if el.tag in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")):
            continue
        el2 = copy.deepcopy(el)
        for blip in el2.iter(qn("a:blip")):
            rid = blip.get(qn("r:embed"))
            part = src.part.related_part(rid)
            blip.set(qn("r:embed"), new.part.relate_to(part, src.part.rels[rid].reltype))
        tree.append(el2)
    return new


def slide_index(slide):
    return [s.slide_id for s in prs.slides].index(slide.slide_id)


def move_after(slide, anchor):
    ids = list(SLD)
    el = ids[slide_index(slide)]
    SLD.remove(el)
    SLD.insert(slide_index(anchor) + 1, el)


def headline(slide, text):
    h = slide.shapes[1]
    fill_tf(h.text_frame, [[(text, {})]], size=fit_size(text), font=F_MED)


def label(slide, text):
    fill_tf(slide.shapes[0].text_frame, [[(text, {})]], size=16, font=F_MED)


def subline(slide, text):
    return tb(slide, X0, 1.4, W, 0.34, [[(text, {})]], size=13, color=SUB, anchor=MSO_ANCHOR.MIDDLE)


# ---------------------------------------------------------------- 표 크기 조정
def set_table(gf, data, col_w, row_h, size=15, head_size=None, bold_cols=(0,), aligns=None, fills=None):
    t = gf.table._tbl
    nr, nc = len(data), len(data[0])
    grid = t.tblGrid
    while len(grid) < nc:
        grid.append(copy.deepcopy(grid[-1]))
    while len(grid) > nc:
        grid.remove(grid[-1])
    for tr in t.tr_lst:
        while len(tr.tc_lst) < nc:
            tr.append(copy.deepcopy(tr.tc_lst[-1]))
        while len(tr.tc_lst) > nc:
            tr.remove(tr.tc_lst[-1])
    while len(t.tr_lst) < nr:
        t.append(copy.deepcopy(t.tr_lst[-1]))
    while len(t.tr_lst) > nr:
        t.remove(t.tr_lst[-1])
    tbl = gf.table
    for i, w in enumerate(col_w):
        tbl.columns[i].width = Inches(w)
    gf.width = Inches(sum(col_w))
    for r in range(nr):
        tbl.rows[r].height = Inches(row_h[r] if isinstance(row_h, (list, tuple)) else row_h)
        for c in range(nc):
            cell = tbl.cell(r, c)
            val = data[r][c]
            runs = val if isinstance(val, list) else [(str(val), {})]
            head = r == 0
            if fills and (r, c) in fills:
                cell.fill.solid()
                cell.fill.fore_color.rgb = rgb(fills[(r, c)])
            fill_tf(cell.text_frame, [runs], size=(head_size or size) if head else size,
                    bold=head or c in bold_cols, color=WHITE if head else TEXT, font=F_MED if head else F_BODY,
                    align=(aligns or {}).get(c, PP_ALIGN.LEFT))


# ---------------------------------------------------------------- 문구 치환 (서식 유지)
def replace_in_paragraph(p, old, new):
    runs = list(p.runs)
    full = "".join(r.text for r in runs)
    i = full.find(old)
    if i < 0:
        return False
    j = i + len(old)
    pos = 0
    first = True
    for r in runs:
        a, b = pos, pos + len(r.text)
        pos = b
        if b <= i or a >= j:
            continue
        s0, s1 = max(i, a) - a, min(j, b) - a
        if first:
            r.text = r.text[:s0] + new + r.text[s1:]
            first = False
        else:
            r.text = r.text[:s0] + r.text[s1:]
    return True


def iter_paragraphs(slide):
    for sh in slide.shapes:
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                yield sh, p
        if sh.has_table:
            for row in sh.table.rows:
                for c in row.cells:
                    for p in c.text_frame.paragraphs:
                        yield sh, p


def edit(slide_head, old, new, tag):
    s = find_slide(slide_head)
    for sh, p in iter_paragraphs(s):
        if replace_in_paragraph(p, old, new):
            LOG["edits"].append((tag, slide_index(s) + 1, old, new or "(삭제)"))
            return
    LOG["edits"].append((tag, None, old, "찾지 못함"))


def delete_paragraph(slide_head, text, tag):
    s = find_slide(slide_head)
    for sh, p in iter_paragraphs(s):
        if p.text.strip() == text:
            p._p.getparent().remove(p._p)
            LOG["edits"].append((tag, slide_index(s) + 1, text, "(삭제)"))
            return
    LOG["edits"].append((tag, None, text, "찾지 못함"))


# =====================================================================
# 기준 장표 (헤드라인으로 찾음, 수정 전 원본 상태에서 복제)
S05 = find_slide("아이라이너로 검증된")
S17 = find_slide("방치된 자산의 재배치")
S19 = find_slide("채널마다 역할을 나누고")
S20 = find_slide("가장 효과가 큰 순서대로")
S21 = find_slide("새 고객과 돌아올 고객을")
S22 = find_slide("고객 리뷰에서 소재를 뽑습니다")
S25 = find_slide("요청하시기 전에, 먼저 가져갑니다")
S26 = find_slide("소재는 감이 아닌")
S18 = find_slide("네이버는 지키고, 메타는 키웁니다")
S24 = find_slide("말씀하신 다섯 가지에 하나씩 답합니다")

v2_text = {}
for i, s in enumerate(prs.slides, 1):
    v2_text[s.slide_id] = "\n".join(p.text for _, p in iter_paragraphs(s))

# ---------------------------------------------------------------- S20A 네이버 검색광고 계정 구조 (S19 복제)
s = duplicate(S19)
label(s, "03 운영 전략")
headline(s, "캠페인을 목적별로 나눠, 예산이 어디서 성과를 내는지 바로 보이게 합니다")
set_table(s.shapes[3], [
    ["캠페인", "유형", "대상", "목적", "입찰 기준"],
    ["브랜드", "파워링크", "나어릴때, 나어릴때 아이라이너, 엣지홀릭", "브랜드 검색 방어", "1위 유지"],
    ["쇼핑_워터프루프", "쇼핑검색", "All New 워터프루프 브러쉬 리퀴드 아이라이너", "핵심 키워드 공략", "광고 상위 구좌 진입"],
    ["쇼핑_펜·브러쉬", "쇼핑검색", "엣지홀릭 펜, 오리지널 브러쉬", "확장", "ROAS 기준 조정"],
    ["용도", "파워링크", "애교살 라이너, 리퀴드 아이라이너, 번지지 않는 아이라이너, 워터프루프 아이라이너", "신규 고객 유입", "CPC 상한 설정"],
], col_w=[2.15, 1.35, 4.6, 2.1, 2.23], row_h=[0.5, 0.55, 0.55, 0.55, 0.78], size=14, head_size=15)
place(s.shapes[3], y=1.8)
bx = s.shapes[4]
place(bx, y=4.95, h=1.8)
shape_text(bx, [
    [("운영 포인트", {"bold": True, "color": LILAC, "size": 14})],
    [("상품명 점검  ", {"bold": True, "color": LILAC}), ("쇼핑검색은 상품명에 '워터프루프', '펜', '브러쉬'가 들어가야 해당 키워드에 노출됩니다. 브랜드스토어 상품명부터 점검합니다", {})],
    [("확장소재  ", {"bold": True, "color": LILAC}), ("리뷰 수, 1+1 혜택, 무료배송 조건을 추가 문구로 노출합니다", {})],
    [("모바일 우선  ", {"bold": True, "color": LILAC}), ("검색의 91%가 모바일이므로 모바일 입찰 가중치를 높게 둡니다", {})],
    [("제외 키워드  ", {"bold": True, "color": LILAC}), ("저가·중고 의도 검색은 효율을 보고 제외합니다", {})],
], size=13, margin=0.3, space_after=2)
S20A = s

# ---------------------------------------------------------------- S20B 브랜드검색 구성안 (S05 복제)
s = duplicate(S05)
sh = list(s.shapes)
remove(sh[0])          # 기존 캡처 사진
remove(sh[3])          # 부제 (원고에 없음)
label(s, "03 운영 전략")
headline(s, "'나어릴때'를 검색하면, 추억 콘텐츠가 아니라 나어릴때가 먼저 보이게 합니다")
sh = list(s.shapes)    # 0 label,1 head,2 page,3 '주력 제품',4~11 카드,12 '판매 채널',13 채널 목록,14 하단 바
image(s, X0, 1.85, 3.3, 4.9, "s20b_brandsearch_wire.png")
cx, cw = 4.05, 5.0
fill_tf(sh[3].text_frame, [[("구성 (모바일 기준)", {})]], size=13, bold=True, color=LILAC, font=F_MED)
place(sh[3], x=cx, y=1.85, w=cw)
items = [("메인 이미지", "엣지홀릭 브러쉬 아이라이너 + \"리뷰 4,901개가 고른 라이너\""),
         ("상품 탭 3개", "오리지널 브러쉬 / 펜 / All New 워터프루프"),
         ("혜택 탭", "1+1 이벤트, 슈퍼적립"),
         ("랜딩", "브랜드스토어")]
for k, (t, d) in enumerate(items):
    card, txt = sh[4 + 2 * k], sh[5 + 2 * k]
    y = 2.25 + k * 0.85
    place(card, x=cx, y=y, w=cw)
    place(txt, x=cx + 0.2, y=y + 0.06, w=cw - 0.4)
    fill_tf(txt.text_frame, [[(t, {"bold": True, "font": F_MED})], [(d, {"size": 12, "color": SUB})]], size=14)
rx = cx + cw + 0.3
fill_tf(sh[12].text_frame, [[("운영", {})]], size=13, bold=True, color=LILAC, font=F_MED)
place(sh[12], x=rx, y=1.85, w=12.88 - rx)
fill_tf(sh[13].text_frame, [
    [("시즌마다 메인 이미지 교체 (11월 슈퍼적립 → 12월 연말 → 1월 신년)", {})],
    [("AI 브리핑보다 위에 노출되어 브랜드명 오해석 문제를 함께 해결합니다", {})],
], size=13, bullet=True, space_after=6)
place(sh[13], x=rx, y=2.25, w=12.88 - rx, h=3.2)
place(sh[14], x=cx, w=12.88 - cx)
shape_text(sh[14], [[("비용  ", {"bold": True}), ("조회수 기준 단가로 산정 〔확인〕브랜드검색 견적", {"bold": True})]], size=15, color=WHITE, margin=0.25)
notes(s, "브랜드 검색량이 크지 않아 단가 부담이 낮은 편입니다. 정확한 금액은 견적 후 안내합니다.")
S20B = s

# ---------------------------------------------------------------- S20C GFA 운영안 (S19 표 + 하단 3칸)
s = duplicate(S19)
label(s, "03 운영 전략")
headline(s, "GFA는 '한 번 본 고객'을 다시 데려오는 데 집중합니다")
set_table(s.shapes[3], [
    ["그룹", "대상", "소재", "시작 예산 비중"],
    ["리타겟팅", "브랜드스토어 방문·장바구니 고객 〔확인〕연동 가능 범위", "혜택 소재 (1+1, 슈퍼적립)", "50%"],
    ["유사 타겟", "구매 고객과 비슷한 사용자", "제품력 소재 (번짐 테스트)", "30%"],
    ["관심사", "30~50대 여성, 뷰티·메이크업 관심", "브랜드 소재 (브러쉬 vs 펜)", "20%"],
], col_w=[2.2, 5.0, 3.33, 1.9], row_h=0.66, size=15, aligns={3: PP_ALIGN.CENTER})
bx = s.shapes[4]
cw3 = (W - 0.6) / 3
three = [("지면", "네이버 피드, 스마트채널 중심으로 시작해 효율 좋은 지면으로 압축"),
         ("소재 형식", "네이티브 이미지 → 이미지 슬라이드 → 동영상 순으로 테스트"),
         ("슈퍼적립 기간", "리타겟팅 비중을 70%까지 올립니다")]
for k, (t, d) in enumerate(three):
    b = bx if k == 0 else clone(s, bx)
    place(b, x=X0 + k * (cw3 + 0.3), y=4.95, w=cw3, h=1.75)
    shape_text(b, [[(t, {"bold": True, "color": LILAC, "size": 15})], [(d, {})]], size=14, margin=0.25, space_after=4, anchor=MSO_ANCHOR.TOP)
    b.text_frame.margin_top = Inches(0.2)
S20C = s

# ---------------------------------------------------------------- S22A 영상 스토리보드 (S22 복제, 4칸 → 5칸)
s = duplicate(S22)
label(s, "03 운영 전략")
headline(s, "15초 안에 '안 번진다'를 눈으로 보여줍니다")
subline(s, "컨셉 1 번짐 테스트 · 9:16 릴스·스토리 · 캠페인 A(신규)")
sh = list(s.shapes)
cards = [sh[3 + 7 * k: 10 + 7 * k] for k in range(4)]
bottom = sh[31]
# 5번째 카드: 4번째 카드 도형 7개 복제
cards.append([clone(s, x) for x in cards[3]])
cuts = [("훅", "0~2초", "\"이거, 진짜 안 번질까?\"", [("화면  ", {"bold": True}), ("손등에 아이라이너로 라인 한 줄", {})], None),
        ("컬러", "2~5초", "\"엣지홀릭 브러쉬 4컬러\"", [("화면  ", {"bold": True}), ("4컬러 스와치 나란히", {})], None),
        ("테스트", "5~9초", "\"물 묻히고, 문질러도\"", [("화면  ", {"bold": True}), ("물을 뿌리고 손가락으로 문지르기", {})], None),
        ("사용", "9~12초", "\"꼬리 라인은 한 번에\"", [("화면  ", {"bold": True}), ("눈꼬리 라인을 한 번에 빼는 클로즈업", {})], None),
        ("증거 + 행동 유도", "12~15초", "\"번지지 않아서 더울 때도 열심히 사용했어요\"", [("화면  ", {"bold": True}), ("리뷰 캡처 + 제품", {})], "★4.7 리뷰 1,154건")]
gap = 0.22
cw5 = (W - 4 * gap) / 5
top = 1.85
for k, (card, (t, sec, cap, scr, extra)) in enumerate(zip(cards, cuts)):
    rect, oval, title, tag, copy_, ln, why = card
    x = X0 + k * (cw5 + gap)
    ix = x + 0.18
    iw = cw5 - 0.36
    place(rect, x=x, y=top, w=cw5, h=4.15)
    place(oval, x=ix, y=top + 0.15)
    fill_tf(oval.text_frame, [[(f"{k + 1}", {"bold": True})]], size=13, color=WHITE, align=PP_ALIGN.CENTER)
    place(title, x=ix + 0.52, y=top + 0.15, w=iw - 0.52, h=0.42)
    fill_tf(title.text_frame, [[(f"컷{k + 1} {t}", {"bold": True})]], size=15 if len(t) > 4 else 16, font=F_MED)
    title.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    place(tag, x=ix, y=top + 0.68, w=iw)
    fill_tf(tag.text_frame, [[(sec, {"bold": True})]], size=12, color=LILAC, align=PP_ALIGN.CENTER)
    image(s, ix, top + 1.1, iw, 0.85, "s22a_storyboard.png")
    place(copy_, x=ix, y=top + 2.05, w=iw, h=1.05)
    paras = [[(cap, {"bold": True, "color": LILAC})]]
    if extra:
        paras.append([(extra, {"bold": True, "size": 12})])
    fill_tf(copy_.text_frame, paras, size=13, font=F_MED)
    place(ln, x=ix, y=top + 3.15, w=iw)
    place(why, x=ix, y=top + 3.22, w=iw, h=0.9)
    wp = [scr]
    if k == 4:
        wp.append([("버튼  ", {"bold": True}), ("1+1로 만나보기", {})])
    fill_tf(why.text_frame, wp, size=12, color=SUB)
place(bottom, y=6.15)
shape_text(bottom, [[("컷3은 실제 촬영 결과로만 사용합니다. 결과가 기대와 다르면 컷3을 '하루 착용 전후' 비교로 바꿉니다.", {})]], size=13, margin=0.3)
notes(s, "기존 제품 컷 + 간단 촬영으로 제작 가능. 제작비 추가 없음.")
S22A = s

# ---------------------------------------------------------------- S22B 배너·카피 세트 (S19 표 + 우측 이미지 2개)
s = duplicate(S19)
label(s, "03 운영 전략")
headline(s, "컨셉마다 메인 문구, 보조 문구, 본문까지 바로 쓸 수 있게 준비했습니다")
set_table(s.shapes[3], [
    ["컨셉", "메인 문구", "보조 문구", "메타 본문 (첫 줄)", "캠페인"],
    ["번짐 테스트", "더울 때도, 하루 종일", "번지지 않는 엣지홀릭 브러쉬", "리뷰 1,154건이 말하는 '안 번지는 라인'", "A 신규"],
    ["라인 시연", "꼬리 라인, 한 번이면 끝", "브러쉬 끝이 살아 있는 아이라이너", "곰손도 한 번에 그려지는 이유", "A 신규"],
    ["브러쉬 vs 펜", "붓처럼? 펜처럼?", "나한테 맞는 라인 고르기", "같은 나어릴때, 다른 라인. 당신은 어느 쪽?", "A·C 테스트"],
    ["브랜드 향수", "나 어릴 때 그 눈매, 지금 다시", "클래식 메이크업 브랜드 나어릴때", "그때 그 또렷함을 지금의 나에게", "A 신규 (40~50대)"],
], col_w=[1.35, 1.85, 1.95, 2.3, 1.15], row_h=[0.45, 0.8, 0.8, 0.8, 0.8], size=13, head_size=13)
place(s.shapes[3], y=1.8)
ix = X0 + 8.6 + 0.3
iw = 12.88 - ix
for k, (f, l) in enumerate([("s22b_banner_1.png", "번짐 테스트"), ("s22b_banner_2.png", "브러쉬 vs 펜")]):
    y = 1.8 + k * 1.95
    image(s, ix, y, iw, 1.6, f)
    tb(s, ix, y + 1.62, iw, 0.26, [[(f"배너 시안 — {l}", {})]], size=11, color=CAP, align=PP_ALIGN.CENTER)
bx = s.shapes[4]
place(bx, y=5.75, h=0.78)
shape_text(bx, [
    [("버튼 문구  ", {"bold": True, "color": LILAC}), ("지금 구매하기 / 1+1 보러 가기 / 내 라인 찾기", {})],
    [("사이즈  ", {"bold": True, "color": LILAC}), ("1:1 (피드), 9:16 (스토리·릴스), 4:5 (피드 세로)", {})],
], size=13, margin=0.3, space_after=2)
tb(s, X0, 6.58, W, 0.3, [[("문구는 방향 예시이며, 화장품 광고 표현 기준 검토 후 확정합니다", {})]], size=12, color=CAP)
notes(s, "디자인팀에서 번짐 테스트·브러쉬 vs 펜 2종만 시안을 만들어 넣으면 충분합니다.")
S22B = s

# ---------------------------------------------------------------- S22C 시즌 혜택 소재 (S17 복제)
s = duplicate(S17)
label(s, "03 운영 전략")
headline(s, "혜택 소재는 '이미 관심 있는 고객'에게만 보여줍니다")
subline(s, "캠페인 B(리타겟팅) · GFA 리타겟팅 전용")
sh = list(s.shapes)
cards = [sh[3 + 5 * k: 8 + 5 * k] for k in range(3)]
benefit = [("11월 슈퍼적립", "\"아이라이너 살 거면, 지금 적립까지\"", "네이버 슈퍼적립 기간 한정 〔확인〕적립률", "브랜드스토어"),
           ("1+1 상시", "\"하나는 파우치에, 하나는 책상에\"", "엣지홀릭 1+1 17,000원", "자사몰"),
           ("12월 연말", "\"선물하기 좋은 라이너 세트\"", "〔확인〕연말 프로모션 구성", "자사몰")]
for k, (rect, num, title, desc, inner) in enumerate(cards):
    t, main, subc, land = benefit[k]
    place(rect, y=1.85, h=3.95)
    place(num, y=2.0)
    place(title, y=2.95)
    fill_tf(title.text_frame, [[(t, {"bold": True})]], size=19, font=F_MED)
    place(desc, y=3.5, h=1.0)
    fill_tf(desc.text_frame, [[(main, {"bold": True, "color": LILAC, "font": F_MED})], [("보조 문구  ", {"bold": True}), (subc, {})]],
            size=14, space_after=6)
    place(inner, y=4.7, h=0.9)
    shape_text(inner, [[("랜딩  ", {"bold": True, "color": LILAC}), (land, {"bold": True, "color": LILAC})]], size=14, margin=0.2)
bx = box(s, X0, 5.98, W, 0.77, fill=LILAC_T)
shape_text(bx, [[("신규 고객에게는 혜택보다 제품력을 먼저 보여주고, 혜택은 재방문 고객에게 보여줘 할인 의존도를 낮춥니다", {"bold": True, "color": LILAC})]], size=15, margin=0.3)
notes(s, "현재 나어릴때 광고는 신규 고객에게도 할인 배너만 노출되는 구조라는 점과 대비됩니다.")
S22C = s

# ---------------------------------------------------------------- S25A 채널 운영 리듬 (S19 복제)
s = duplicate(S19)
label(s, "04 운영 방식")
headline(s, "매일, 매주, 매월 할 일을 정해두고 운영합니다")
set_table(s.shapes[3], [
    ["주기", "네이버 SA·GFA", "메타", "인스타그램·리포트"],
    ["매일", "예산 소진 속도, 순위 이상 확인", "예산 소진, 구매 이상치 확인", "단톡방 요청 응대"],
    ["매주", "입찰 조정, 키워드 추가·제외", "소재 성과 판별, 신규 소재 투입", "주간 리포트 + 이번 주 제안, 피드 2회 업로드"],
    ["매월", "키워드 구조 정리, 상품명 점검", "캠페인 예산 재배분, 타겟 점검", "월간 리뷰 미팅, 다음 달 시즌 제안"],
], col_w=[1.5, 3.55, 3.55, 3.83], row_h=[0.66, 0.85, 0.85, 0.85], size=15)
shape_text(s.shapes[4], [[("첫 달은 성과 기준선을 만드는 달입니다. 2개월 차부터 목표 ROAS를 함께 정해 관리합니다", {"bold": True, "color": LILAC})]], size=16, margin=0.3)
S25A = s

# ---------------------------------------------------------------- S21 메타 계정 구조 (교체: 캠페인 2 → 3, 기존 디자인 유지)
s = S21
headline(s, "새 고객, 돌아올 고객, 테스트를 나눠서 운영합니다")
for x in list(s.shapes)[3:]:
    remove(x)
camps = [("캠페인 A — 신규 고객", "예산 60%", LILAC, [
            ("방식", "Advantage+ 쇼핑 캠페인 또는 수동 캠페인 비교 후 선택"),
            ("타겟", "30~50대 여성 중심, 넓은 타겟에서 메타 AI가 구매 가능 고객 탐색"),
            ("소재", "제품력 중심 (번짐 테스트, 라인 시연, 브랜드 향수)"),
            ("지표", "CPA, 구매 ROAS")]),
         ("캠페인 B — 리타겟팅", "예산 25%", LILAC_M, [
            ("방식", "카탈로그 광고 (카페24 상품 정보 연동, 본 상품을 다시 보여줌)"),
            ("타겟", "자사몰 방문 30일, 장바구니 7일, 구매자 제외"),
            ("소재", "혜택 중심 (1+1, 슈퍼적립, 기간 한정)"),
            ("지표", "구매 ROAS")]),
         ("캠페인 C — 소재 테스트", "예산 15%", LILAC_M, [
            ("", "새 소재를 같은 조건으로 먼저 검증"),
            ("", "이긴 소재만 캠페인 A로 이동 (S26 위닝 소재 운영과 연결)")])]
cw3 = (W - 0.5) / 3
for k, (t, budget, col, rows) in enumerate(camps):
    x = X0 + k * (cw3 + 0.25)
    box(s, x, 1.75, cw3, 3.3, fill=WHITE, line=col, line_w=1.25)
    tb(s, x + 0.22, 1.88, cw3 - 1.5, 0.4, [[(t, {"bold": True, "color": LILAC})]], size=15, font=F_MED, anchor=MSO_ANCHOR.MIDDLE)
    chip(s, x + cw3 - 1.2, 1.93, 1.0, budget, fill=LILAC if k == 0 else LILAC_T, color=WHITE if k == 0 else LILAC, size=11)
    y = 2.45
    for key, val in rows:
        if key:
            chip(s, x + 0.22, y + 0.02, 0.6, key, size=11, h=0.28)
            tb(s, x + 0.95, y, cw3 - 1.15, 0.62, [[(val, {})]], size=12)
            y += 0.64
        else:
            tb(s, x + 0.22, y, cw3 - 0.44, 0.62, [[(val, {})]], size=13, bullet=True)
            y += 0.7
hw = (W - 0.25) / 2
b = box(s, X0, 5.2, hw, 1.55, fill=LILAC_T)
shape_text(b, [
    [("확장 옵션 — 파트너십 광고", {"bold": True, "color": LILAC})],
    [("실구매 리뷰어·뷰티 크리에이터 게시물을 광고로 집행", {"size": 12})],
    [("경쟁사(노에트 등)도 활용 중인 방식 〔확인〕크리에이터 섭외비 별도 여부", {"size": 12})],
], size=14, margin=0.25, space_after=3, anchor=MSO_ANCHOR.MIDDLE)
b = box(s, X0 + hw + 0.25, 5.2, hw, 1.55, fill=PEACH_T)
shape_text(b, [
    [("측정 기반", {"bold": True, "color": LILAC})],
    [("메타 픽셀 구매 이벤트 점검 + 전환 API 연동", {"size": 12})],
    [("광고 링크 UTM 규칙 통일", {"size": 12})],
    [("GA4 설치로 채널 간 성과 비교", {"size": 12})],
], size=14, margin=0.25, space_after=2, anchor=MSO_ANCHOR.MIDDLE)

# ---------------------------------------------------------------- S26 위닝 소재 운영 (보강: 하단 규칙 박스)
s = S26
sh = list(s.shapes)  # 0 label 1 head 2 sub 3 page 4~ 카드/화살표, 18 사례 박스
for k in range(3):
    base = 4 + k * 5
    rect, num, title, desc = sh[base], sh[base + 1], sh[base + 2], sh[base + 3]
    place(rect, y=1.85, h=2.45)
    place(num, y=1.98)
    place(title, y=2.75)
    place(desc, y=3.25, h=1.0)
    if k < 2:
        place(sh[base + 4], y=2.88)
case = sh[18]
place(case, y=4.45, h=0.75)
rules = box(s, X0, 5.35, W, 1.4, fill=WHITE, line=LILAC_M, line_w=1.25)
tb(s, X0 + 0.3, 5.47, 3, 0.36, [[("소재 운영 규칙", {"bold": True, "color": LILAC})]], size=15, font=F_MED, anchor=MSO_ANCHOR.MIDDLE)
rl = ["광고 세트마다 소재 3~6개 유지", "같은 사람에게 3회 이상 노출되고 클릭률이 떨어지면 교체",
      "2주마다 신규 소재 투입", "학습 기간(집행 초기)에는 잦은 수정 없이 지켜봅니다"]
for k, t in enumerate(rl):
    tb(s, X0 + 0.3 + (k % 2) * 6.1, 5.9 + (k // 2) * 0.38, 5.9, 0.36, [[(t, {})]], size=13, bullet=True)

# ---------------------------------------------------------------- 간지 소제목 한 줄씩 추가
def add_section_line(sec_slide, after_text, new_text):
    for sh_ in sec_slide.shapes:
        if not sh_.has_text_frame:
            continue
        for p in sh_.text_frame.paragraphs:
            if p.text.strip() == after_text:
                np_ = copy.deepcopy(p._p)
                p._p.addnext(np_)
                runs = np_.findall(qn("a:r"))
                for r in runs[1:]:
                    np_.remove(r)
                runs[0].find(qn("a:t")).text = new_text
                return True
    raise KeyError(after_text)


def section_lines(sec_slide):
    for sh_ in sec_slide.shapes:
        if sh_.has_text_frame and len(sh_.text_frame.paragraphs) >= 2 and sh_.top > Inches(4.8):
            return [p.text for p in sh_.text_frame.paragraphs]
    return []


l18 = section_lines(S18)
add_section_line(S18, l18[1], "소재 예시")
l24 = section_lines(S24)
add_section_line(S24, l24[0], "채널 운영 리듬")

# ---------------------------------------------------------------- 장표 순서
move_after(S20A, S20)
move_after(S20B, S20A)
move_after(S20C, S20B)
move_after(S22A, S22)
move_after(S22B, S22A)
move_after(S22C, S22B)
move_after(S25A, S25)

# ---------------------------------------------------------------- 표현 수정 (승인: 전부)
edit("10월 대행 이관 즉시 고효율", "10월 대행 이관 즉시", "11월 대행 이관 즉시", "1 S15")
edit("10월 대행 이관 즉시 GA4", "10월 대행 이관 즉시", "11월 대행 이관 즉시", "2 S16")
edit("광고주 요청사항 핵심 요약", "월 40건+", "월 20건 내외, 최대 50건", "3 S06")
edit("광고주 요청사항 핵심 요약", "전담 TF", "전담 담당자", "4 S06")
delete_paragraph("소재 제작 지원 안내", "제작 소재가 너무 많아지면 외주에 따른 실비가 발생할 수 있습니다.", "5 S27")
edit("유입 방어선 붕괴", "유입 방어선 붕괴", "광고 구좌 공백", "6 S12")
edit("경쟁사에 전면 노출 방치 중", "경쟁사에 전면 노출 방치 중", "경쟁사가 차지하고 있음", "7 S12")
edit("통합 가격 통제(MAP 관리) 시급", "▶ 통합 가격 통제(MAP 관리) 시급", "▶ 판매처 가격 정책 점검 권장", "8 S13")
edit("강력히 제안합니다", "강력히 제안합니다", "제안드립니다", "9 S13")
edit("인스타그램 피드 방치 및", "피드 방치", "피드 공백", "10 S16")
edit("\"새로운 투자가 아닌", "\"새로운", "새로운", "11 S17")

# ---------------------------------------------------------------- 페이지 번호 다시 매김
for i, sl in enumerate(prs.slides, 1):
    for sh_ in sl.shapes:
        if sh_.has_text_frame and abs(sh_.left - Inches(12.2)) < Inches(0.05) and abs(sh_.top - Inches(7.0)) < Inches(0.05) \
                and sh_.text_frame.text.strip().isdigit():
            r = sh_.text_frame.paragraphs[0].runs[0]
            r.text = str(i)

OUT.parent.mkdir(exist_ok=True)
prs.save(str(OUT))

# ---------------------------------------------------------------- 보고용 정보
NEW = {S20A.slide_id: "S20A 네이버 검색광고 계정 구조", S20B.slide_id: "S20B 브랜드검색 구성안", S20C.slide_id: "S20C GFA 운영안",
       S22A.slide_id: "S22A 소재 예시 ① 영상 스토리보드", S22B.slide_id: "S22B 소재 예시 ② 배너·카피 세트",
       S22C.slide_id: "S22C 소재 예시 ③ 시즌 혜택 소재", S25A.slide_id: "S25A 채널 운영 리듬"}
MOD = {S21.slide_id: "S21 메타 계정 구조 (교체)", S26.slide_id: "S26 위닝 소재 운영 (보강)",
       S18.slide_id: "S18 간지 03 (소제목 추가)", S24.slide_id: "S24 간지 04 (소제목 추가)"}
print("saved", OUT, "slides", len(prs.slides))
for i, sl in enumerate(prs.slides, 1):
    tag = NEW.get(sl.slide_id) or MOD.get(sl.slide_id)
    if tag:
        print(f"  {i:2d}  {tag}")
print("edits:")
for e in LOG["edits"]:
    print("  ", e)
print("missing:", sorted(LOG["missing"]))
print("checks:", LOG["checks"])
