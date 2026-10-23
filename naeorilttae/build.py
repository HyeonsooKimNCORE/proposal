"""나어릴때 디지털 광고 운영 제안서 생성 스크립트 (작업지시서.md 2단계)

실행: python3 build.py  →  output/나어릴때_제안서_v1.pptx
- 템플릿: template/template.pptx (Cover / Section / Content 레이아웃 사용, 샘플 슬라이드 삭제)
- 원고: 나어릴때_제안서_원고.md (문구·수치 기준)
- 색: 템플릿 표지의 라일락 붓터치·복숭아 배경에서 뽑은 색만 사용
- 폰트: Pretendard / Pretendard Medium / Pretendard Black (템플릿 사용 폰트)
- 없는 이미지는 회색 박스에 [이미지 필요: 파일명] 표시
"""
import copy
import csv
import re
from collections import Counter
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).parent
ASSETS = ROOT / "assets"
OUT = ROOT / "output" / "나어릴때_제안서_v1.pptx"

# ---------- 색 (템플릿 표지 이미지에서 추출) ----------
TEXT = "211D1E"      # 나어릴때 로고 글자색
SUB = "5E5759"
CAP = "8F8789"
LINE = "BFBFBF"      # 템플릿 본문 가로선
WHITE = "FFFFFF"
LILAC = "7B2FB5"     # 라일락 붓터치(진한 부분)
LILAC_M = "C9A3E3"   # 라일락 중간
LILAC_T = "F1E4F8"   # 라일락 연한 배경
PEACH = "F6D2A8"     # 복숭아(진한 쪽)
PEACH_T = "FEF2DA"   # 표지 배경 복숭아
IMG_BG = "E6E1DE"
CHECK = "FFF176"     # 〔확인〕 노란 하이라이트

F_BODY, F_MED, F_BLACK = "Pretendard", "Pretendard Medium", "Pretendard Black"

# 본문 영역 (템플릿 Content 레이아웃의 두 가로선 기준: x 0.45~12.88, 위 0.58 / 아래 6.95)
X0, X1 = 0.45, 12.88
W = X1 - X0
BODY_Y = 1.95
BODY_END = 6.78

prs = Presentation(str(ROOT / "template" / "template.pptx"))

# 템플릿 샘플 슬라이드 삭제
sld_ids = prs.slides._sldIdLst
for sld in list(sld_ids):
    prs.part.drop_rel(sld.rId)
    sld_ids.remove(sld)

LAYOUT = {l.name: l for l in prs.slide_layouts}


def rgb(h):
    return RGBColor.from_string(h)


# ---------- 텍스트 ----------
CHECK_MARK = "〔확인〕"


def split_check(text):
    """〔확인〕 구간을 분리. 숫자로 시작하면 다음 공백까지, 아니면 ' / ' 또는 ')' 앞까지(없으면 끝까지)."""
    out, i = [], 0
    while True:
        j = text.find(CHECK_MARK, i)
        if j < 0:
            out.append((text[i:], False))
            break
        out.append((text[i:j], False))
        k = j + len(CHECK_MARK)
        if k < len(text) and text[k].isdigit():
            m = re.search(r"\s", text[k:])
            e = k + m.start() if m else len(text)
        else:
            cands = [p for p in (text.find(" / ", k), text.find(")", k)) if p >= 0]
            e = min(cands) if cands else len(text)
        out.append((text[j:e], True))
        i = e
    return [(t, c) for t, c in out if t]


def set_run_font(run, size, bold=False, color=TEXT, font=F_BODY, highlight=False):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = rgb(color)
    f.name = font
    rPr = run._r.get_or_add_rPr()
    latin = rPr.find(qn("a:latin"))
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
            latin.addnext(el) if tag == "a:ea" else rPr.find(qn("a:ea")).addnext(el)
        el.set("typeface", font)
    if highlight:
        hl = etree.Element(qn("a:highlight"))
        etree.SubElement(hl, qn("a:srgbClr")).set("val", CHECK)
        rPr.find(qn("a:solidFill")).addnext(hl)


def fill_tf(tf, paras, size=14, color=TEXT, font=F_BODY, bold=False, align=PP_ALIGN.LEFT,
            spacing=1.15, space_after=0, bullet=False):
    """paras: str | list[str | list[(text, opts)]]. opts: size,bold,color,font"""
    if isinstance(paras, str):
        paras = [paras]
    tf.clear()
    for pi, para in enumerate(paras):
        p = tf.paragraphs[0] if pi == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        if space_after:
            p.space_after = Pt(space_after)
        pPr = p._p.get_or_add_pPr()
        pPr.set("latinLnBrk", "0")
        pPr.set("eaLnBrk", "1")
        runs = [(para, {})] if isinstance(para, str) else para
        is_bullet = bullet
        for text, opts in runs:
            if opts.get("bullet"):
                is_bullet = True
            for seg, chk in split_check(text):
                r = p.add_run()
                r.text = seg
                set_run_font(r, opts.get("size", size), opts.get("bold", bold), opts.get("color", color),
                             opts.get("font", font), highlight=chk)
        if is_bullet:
            sz = runs[0][1].get("size", size)
            pPr.set("marL", str(int(Pt(sz * 1.0))))
            pPr.set("indent", str(-int(Pt(sz * 1.0))))
            bu = etree.SubElement(pPr, qn("a:buChar"))
            bu.set("char", "•")


def tb(slide, x, y, w, h, paras, size=14, anchor=MSO_ANCHOR.TOP, name=None, **kw):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", 0)
    tf.vertical_anchor = anchor
    fill_tf(tf, paras, size=size, **kw)
    if name:
        box.name = name
    return box


# ---------- 도형 ----------
def _strip_style(shape):
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


def box(slide, x, y, w, h, fill=WHITE, line=None, radius=0.08, shape=MSO_SHAPE.ROUNDED_RECTANGLE, line_w=1.0, dash=None, name=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    _strip_style(s)
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(line_w)
        if dash:
            s.line.dash_style = dash
    else:
        s.line.fill.background()
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = min(radius / min(w, h), 0.5)
    if name:
        s.name = name
    return s


def shape_text(s, paras, size=14, anchor=MSO_ANCHOR.MIDDLE, margin=0.12, **kw):
    tf = s.text_frame
    tf.word_wrap = True
    for side in ("left", "right"):
        setattr(tf, f"margin_{side}", Inches(margin))
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    tf.vertical_anchor = anchor
    fill_tf(tf, paras, size=size, **kw)


def badge(slide, x, y, text, d=0.42, fill=LILAC, color=WHITE, size=13):
    s = box(slide, x, y, d, d, fill=fill, shape=MSO_SHAPE.OVAL)
    shape_text(s, [[(text, {"bold": True})]], size=size, color=color, align=PP_ALIGN.CENTER, margin=0)
    return s


def chip(slide, x, y, w, text, fill=LILAC_T, color=LILAC, size=12, h=0.32, bold=True):
    s = box(slide, x, y, w, h, fill=fill, radius=h / 2)
    shape_text(s, [[(text, {"bold": bold})]], size=size, color=color, align=PP_ALIGN.CENTER, margin=0.05)
    return s


def hline(slide, x, y, w, color=LINE, width=0.75):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(x + w), Inches(y))
    _strip_style(c)
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    return c


def arrow(slide, x, y, w, h, color=LILAC_M, direction="right"):
    shp = {"right": MSO_SHAPE.RIGHT_ARROW, "down": MSO_SHAPE.DOWN_ARROW}[direction]
    return box(slide, x, y, w, h, fill=color, shape=shp)


def image(slide, x, y, w, h, fname, fit="contain"):
    """assets/에 파일이 있으면 넣고, 없으면 회색 자리 표시"""
    p = ASSETS / fname
    if p.exists():
        from PIL import Image as PImage
        iw, ih = PImage.open(p).size
        r = min(w / iw, h / ih)
        pw, ph = iw * r, ih * r
        pic = slide.shapes.add_picture(str(p), Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2), Inches(pw), Inches(ph))
        pic.name = f"IMG_{fname}"
        return pic
    MISSING.add(fname)
    s = box(slide, x, y, w, h, fill=IMG_BG, radius=0.05, name=f"IMG_{fname}")
    shape_text(s, [[(f"[이미지 필요: {fname}]", {})]], size=12, color=SUB, align=PP_ALIGN.CENTER)
    return s


def notes(slide, text):
    if text and text != "-":
        slide.notes_slide.notes_text_frame.text = text


# ---------- 표 ----------
NO_STYLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"


def _cell_border(cell, color=LINE, width=0.75):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        old = tcPr.find(qn(tag))
        if old is not None:
            tcPr.remove(old)
        ln = etree.SubElement(tcPr, qn(tag), w=str(int(Pt(width))))
        sf = etree.SubElement(ln, qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr")).set("val", color)
    # 테두리는 채우기보다 앞에 와야 함
    fill = tcPr.find(qn("a:solidFill"))
    if fill is not None:
        tcPr.remove(fill)
        tcPr.append(fill)


def table(slide, x, y, col_w, rows, row_h=0.42, size=13, header_fill=LILAC, fills=None, bold_cols=(), aligns=None):
    """rows[0] = 머리행. fills: {(r,c) or r: hex}"""
    nr, nc = len(rows), len(rows[0])
    gf = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(sum(col_w)), Inches(row_h * nr))
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "0")
    tblPr.set("bandRow", "0")
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is None:
        sid = etree.SubElement(tblPr, qn("a:tableStyleId"))
    sid.text = NO_STYLE
    for i, w in enumerate(col_w):
        tbl.columns[i].width = Inches(w)
    for r in range(nr):
        tbl.rows[r].height = Inches(row_h)
        for c in range(nc):
            cell = tbl.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(0.1)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            head = r == 0
            fill = header_fill if head else WHITE
            if fills:
                fill = fills.get((r, c), fills.get(r, fill))
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(fill)
            val = rows[r][c]
            runs = val if isinstance(val, list) else [(str(val), {})]
            al = (aligns or {}).get(c, PP_ALIGN.LEFT)
            fill_tf(cell.text_frame, [runs], size=size, bold=head or c in bold_cols,
                    color=WHITE if head else TEXT, font=F_MED if head else F_BODY, align=al)
            _cell_border(cell)
    return gf


# ---------- 차트 ----------
def chart_font(chart, size=11):
    chart.font.name = F_BODY
    chart.font.size = Pt(size)
    chart.font.color.rgb = rgb(SUB)


def doughnut(slide, x, y, w, h, title, cats, vals, colors):
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series(title, vals)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    chart_font(ch, 12)
    ch.has_title = True
    ch.chart_title.text_frame.text = title
    fill_tf(ch.chart_title.text_frame, [[(title, {"bold": True})]], size=14, font=F_MED)
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch.legend.include_in_layout = False
    pl = ch.plots[0]
    pl.has_data_labels = True
    dl = pl.data_labels
    dl.number_format = '0"%"'
    dl.number_format_is_linked = False
    dl.show_value = True
    dl.font.size = Pt(13)
    dl.font.bold = True
    for i, col in enumerate(colors):
        pt = pl.series[0].points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = rgb(col)
        pt.data_label.font.color.rgb = rgb(WHITE if col == LILAC else TEXT)
        pt.data_label.font.bold = True
        pt.data_label.font.size = Pt(13)
    hole = pl._element.find(qn("c:holeSize"))
    if hole is not None:
        hole.set("val", "55")
    return ch


# =====================================================================
MISSING = set()
CHECKS = []


def content(section, headline, sub=None, source=None, page=None, hsize=None):
    s = prs.slides.add_slide(LAYOUT["Content"])
    tb(s, 0.4, 0.12, 8, 0.42, [[(section, {})]], size=16, font=F_MED, anchor=MSO_ANCHOR.MIDDLE, color=TEXT)
    # 헤드라인: 원고 그대로, 24pt에서 시작해 한 줄에 맞게 최소 20pt까지 줄임
    size = hsize or fit_size(headline, W, 24, 20)
    tb(s, X0, 0.78, W, 0.62, [[(headline, {})]], size=size, font=F_MED, anchor=MSO_ANCHOR.MIDDLE)
    if sub:
        tb(s, X0, 1.4, W, 0.34, [[(sub, {})]], size=13, color=SUB, anchor=MSO_ANCHOR.MIDDLE)
    if source:
        tb(s, X0, 7.0, 10, 0.28, [[(source, {})]], size=10, color=CAP, anchor=MSO_ANCHOR.MIDDLE)
    tb(s, 12.2, 7.0, 0.68, 0.28, [[(str(len(prs.slides._sldIdLst)), {})]], size=10, color=CAP, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    return s


def text_width(text, size):
    w = 0
    for ch in text:
        if "가" <= ch <= "힣" or ch in "〔〕":
            w += 0.95
        elif ch == " ":
            w += 0.28
        else:
            w += 0.58
    return w * size / 72


def fit_size(text, width, start, minimum):
    s = start
    while s > minimum and text_width(text, s) > width - 0.1:
        s -= 1
    return s


def section_slide(num, title, sub, items):
    s = prs.slides.add_slide(LAYOUT["Section"])
    tb(s, 1.0, 1.9, 4, 1.2, [[(num, {})]], size=72, font=F_BLACK, color=LILAC, anchor=MSO_ANCHOR.BOTTOM)
    tb(s, 1.0, 3.2, 10, 0.8, [[(title, {})]], size=40, font=F_MED, anchor=MSO_ANCHOR.MIDDLE)
    tb(s, 1.0, 4.05, 10, 0.45, [[(sub, {})]], size=18, color=SUB, anchor=MSO_ANCHOR.MIDDLE)
    hline(s, 1.0, 4.75, 5.5, color=LILAC_M, width=1)
    tb(s, 1.0, 4.95, 10, 1.4, [[(t, {})] for t in items], size=14, color=SUB, space_after=4)
    return s


# =====================================================================
# S01 표지
s = prs.slides.add_slide(LAYOUT["Cover"])
s.shapes.add_picture(str(ASSETS / "logo_naeorilttae.png"), Inches(1.16), Inches(1.45), Inches(2.6), Inches(0.72))
tb(s, 1.16, 2.35, 7.0, 1.6, [[("나어릴때 디지털", {})], [("광고 운영 제안", {})]], size=40, font=F_BLACK, spacing=1.1)
tb(s, 1.16, 4.05, 7.0, 0.45, [[("이미 가진 강점을, 광고가 제대로 쓰게 만들겠습니다", {})]], size=20, font=F_MED, color=SUB)
tb(s, 1.16, 5.0, 5, 0.95, [[("나어릴때 귀중", {})], [("2026.10.30", {})]], size=16, font=F_MED, spacing=1.2)
s.shapes.add_picture(str(ASSETS / "logo_nicecore.png"), Inches(1.16), Inches(6.1), Inches(1.7), Inches(0.25))

# S02 목차
s = content("", "Contents")
toc = [
    ("01", "현황 이해", "말씀하신 내용과 브랜드 현황"),
    ("02", "사전 진단", "시장, 네이버, 메타, 인스타그램"),
    ("03", "운영 전략", "채널별 역할과 실행안"),
    ("04", "운영 방식", "역제안, 소재, 인스타그램, 리포트, 소통"),
    ("05", "비용과 일정", ""),
]
for i, (n, t, d) in enumerate(toc):
    y = 1.75 + i * 0.98
    tb(s, 1.0, y, 1.1, 0.7, [[(n, {})]], size=30, font=F_BLACK, color=LILAC, anchor=MSO_ANCHOR.MIDDLE)
    tb(s, 2.2, y, 3.2, 0.7, [[(t, {})]], size=22, font=F_MED, anchor=MSO_ANCHOR.MIDDLE)
    if d:
        tb(s, 5.4, y, 7.0, 0.7, [[(d, {})]], size=15, color=SUB, anchor=MSO_ANCHOR.MIDDLE)
    hline(s, 1.0, y + 0.84, 11.4)

# S03 제안 요약
s = content("제안 요약", "세 가지 공백을 채우면, 11월부터 달라집니다", "외부 공개 데이터만으로 확인한 진단 결과와 해결 방향")
cards = [
    ("네이버 핵심 키워드에 나어릴때 광고가 없습니다", "'워터프루프·펜 아이라이너' 쇼핑검색·파워링크 노출 0, 브랜드검색 미집행", "검색 자리 되찾기 — 워터프루프 쇼핑검색 + 브랜드검색"),
    ("고객은 '안 번진다'고 말하는데, 광고는 할인만 말합니다", "메타 광고 5개 모두 할로윈 할인 배너, 제품 시연·영상 없음", "리뷰 기반 소재 — 번짐 테스트, 라인 시연, 브러쉬 vs 펜"),
    ("광고는 돌고 있는데, 인스타그램은 멈춰 있습니다", "9/12 이후 게시물 없음 (40일 이상 공백)", "광고와 연결된 계정 운영 — 주 2회, 광고 시즌과 피드 연동"),
]
cw = (W - 0.6) / 3
for i, (f, g, sol) in enumerate(cards):
    x = X0 + i * (cw + 0.3)
    box(s, x, BODY_Y, cw, 2.35, fill=WHITE, line=LINE)
    chip(s, x + 0.25, BODY_Y + 0.22, 1.0, f"발견 {i + 1}")
    tb(s, x + 0.25, BODY_Y + 0.65, cw - 0.5, 0.8, [[(f, {"bold": True})]], size=15, font=F_MED)
    tb(s, x + 0.25, BODY_Y + 1.45, cw - 0.5, 0.8, [[("근거  ", {"bold": True, "color": LILAC}), (g, {})]], size=12, color=SUB)
    arrow(s, x + cw / 2 - 0.2, BODY_Y + 2.45, 0.4, 0.4, direction="down")
    sb = box(s, x, BODY_Y + 2.95, cw, 1.85, fill=LILAC_T)
    chip(s, x + 0.25, BODY_Y + 3.17, 1.0, "해결", fill=LILAC, color=WHITE)
    head, tail = sol.split(" — ")
    tb(s, x + 0.25, BODY_Y + 3.6, cw - 0.5, 1.15, [[(head, {"bold": True, "color": LILAC})], [(tail, {})]], size=15, font=F_MED, space_after=2)
notes(s, "이 한 장만 보셔도 제안 전체가 이해되도록 구성했습니다.")

# S04 간지
section_slide("01", "현황 이해", "말씀해주신 내용과 브랜드 현황을 정리했습니다", ["브랜드 이해", "요청 사항 이해"])

# S05 브랜드 이해
s = content("01 현황 이해", "나어릴때는 '아이라이너로 검증된' 브랜드입니다", "자사몰 naeorilttae.com 기준 (2026.10.23)")
tb(s, X0, BODY_Y - 0.05, 6, 0.3, [[("주력 제품", {"bold": True, "color": LILAC})]], size=13, font=F_MED)
prods = [
    ("오리지널 엣지홀릭 브러쉬 아이라이너", "리뷰 1,154건 · 4.7점 | 8,500원 (1+1 17,000원)"),
    ("엣지홀릭 펜 아이라이너", "리뷰 778건 · 4.7점 | 8,500원"),
    ("All New 워터프루프 엣지홀릭 브러쉬 리퀴드 아이라이너", "8,500원 (1+1 14,900원)"),
    ("확장 제품", "글로시 리퀴드 립스틱(리뷰 301건), 알로에 수딩젤(리뷰 394건)"),
]
for i, (n, d) in enumerate(prods):
    y = BODY_Y + 0.35 + i * 0.85
    box(s, X0, y, 6.2, 0.75, fill=WHITE if i < 3 else PEACH_T, line=LINE if i < 3 else None)
    tb(s, X0 + 0.2, y + 0.06, 5.8, 0.65, [[(n, {"bold": True, "font": F_MED})], [(d, {"size": 12, "color": SUB})]], size=14, anchor=MSO_ANCHOR.MIDDLE)
rx = X0 + 6.5
tb(s, rx, BODY_Y - 0.05, 5, 0.3, [[("판매 채널", {"bold": True, "color": LILAC})]], size=13, font=F_MED)
tb(s, rx, BODY_Y + 0.35, W - 6.5, 1.0, [
    [("자사몰: naeorilttae.com (카페24)", {})],
    [("네이버: 브랜드스토어 brand.naver.com/naeorilttae", {})],
    [("네이버 가격비교 카탈로그 리뷰 4,901건 (엣지홀릭 브러쉬)", {})],
], size=13, bullet=True, space_after=2)
image(s, rx, BODY_Y + 1.45, W - 6.5, 2.3, "s05_mall_main.png")
b = box(s, X0, 5.95, W, 0.75, fill=LILAC)
shape_text(b, [[("매출의 중심은 아이라이너, 립·스킨은 확장 라인 → 광고는 '아이라이너 히어로'로 집중합니다", {"bold": True})]], size=15, color=WHITE, margin=0.25)
notes(s, "색조 브랜드 전체가 아니라, 아이라이너 1위 상품을 중심으로 광고를 설계하겠다는 출발점입니다.")

# S06 요청 사항 이해
s = content("01 현황 이해", "말씀하신 다섯 가지, 이렇게 이해했습니다")
reqs = [("소통", "업무시간 안에 바로 응대가 되는 담당자"), ("분석", "숫자 나열이 아니라 '왜'와 '다음 액션'이 있는 리포트"),
        ("역제안", "요청하기 전에 먼저 기회를 가져오는 대행사"), ("소재 제작", "추가 비용 걱정 없이 충분한 물량"),
        ("인스타그램", "광고와 톤이 맞는 계정 운영")]
for i, (t, d) in enumerate(reqs):
    y = 1.65 + i * 0.74
    box(s, X0, y, W, 0.62, fill=WHITE, line=LINE)
    tb(s, X0 + 0.25, y, 0.8, 0.62, [[(f"0{i + 1}", {})]], size=20, font=F_BLACK, color=LILAC, anchor=MSO_ANCHOR.MIDDLE)
    tb(s, X0 + 1.1, y, 2.2, 0.62, [[(t, {"bold": True})]], size=17, font=F_MED, anchor=MSO_ANCHOR.MIDDLE)
    tb(s, X0 + 3.3, y, W - 3.5, 0.62, [[(d, {})]], size=15, color=SUB, anchor=MSO_ANCHOR.MIDDLE)
b = box(s, X0, 5.45, W, 1.25, fill=LILAC_T)
shape_text(b, [[("대행사를 바꾸시려는 이유는 '성과'보다 '같이 일하는 방식'이라고 이해했습니다.", {"bold": True, "color": LILAC})],
               [("이 제안서는 그 다섯 가지에 하나씩 답하는 구조입니다.", {})]], size=15, margin=0.3, space_after=3)
notes(s, "S25~S30에서 다섯 가지에 하나씩 답합니다.")

# S07 간지
section_slide("02", "사전 진단", "광고 계정을 보기 전에, 밖에서 보이는 것부터 확인했습니다",
              ["시장 · 네이버 · 메타 · 인스타그램과 자사몰", "진단 종합"])

# S08 진단 범위
s = content("02 사전 진단", "6개 영역을 직접 확인했습니다", "조사일 2026.10.23 · 외부 공개 데이터 기준")
areas = [
    ("네이버 데이터랩", "아이라이너 카테고리 고객·검색어 추이 (최근 1년)"),
    ("네이버 통합검색", "'나어릴때' 검색 결과 (PC·모바일)"),
    ("네이버 쇼핑", "'워터프루프 아이라이너', '펜 아이라이너' 노출 현황"),
    ("메타 광고 라이브러리", "나어릴때 및 아이라이너 경쟁 광고 약 110개"),
    ("인스타그램", "@naeorilttae_official 운영 현황"),
    ("자사몰", "상품 구성, 리뷰, 광고 추적 설정"),
]
cw = (W - 0.6) / 3
for i, (t, d) in enumerate(areas):
    x, y = X0 + (i % 3) * (cw + 0.3), BODY_Y + (i // 3) * 1.85
    box(s, x, y, cw, 1.65, fill=WHITE, line=LINE)
    badge(s, x + 0.25, y + 0.25, str(i + 1))
    tb(s, x + 0.8, y + 0.25, cw - 1.0, 0.42, [[(t, {"bold": True})]], size=16, font=F_MED, anchor=MSO_ANCHOR.MIDDLE)
    tb(s, x + 0.25, y + 0.85, cw - 0.5, 0.7, [[(d, {})]], size=13, color=SUB)
tb(s, X0, 5.75, W, 0.6, [[("광고 계정 성과 데이터는 포함하지 않았습니다. 계정 권한을 주시면 〔확인〕2주 안에 정밀 진단 리포트를 드립니다.", {})]], size=13, color=SUB)

# S09 시장 진단 ① 고객
s = content("02 사전 진단", "아이라이너는 40~50대 여성이 모바일로 찾는 시장입니다",
            "네이버 데이터랩 쇼핑인사이트 · 화장품/미용 > 색조메이크업 > 아이라이너 · 2025.10.22~2026.10.22",
            source="출처: 네이버 데이터랩 (카테고리 전체 기준, 나어릴때 고객 데이터 아님)")
for i, (t, cats, vals) in enumerate([("기기", ["모바일", "PC"], [91, 9]), ("성별", ["여성", "남성"], [92, 8])]):
    box(s, X0 + i * 3.25, BODY_Y, 3.05, 3.2, fill=WHITE, line=LINE)
    doughnut(s, X0 + i * 3.25 + 0.1, BODY_Y + 0.05, 2.85, 3.1, t, cats, vals, [LILAC, PEACH])
box(s, X0 + 6.5, BODY_Y, W - 6.5, 3.2, fill=WHITE, line=LINE)
tb(s, X0 + 6.75, BODY_Y + 0.12, 4, 0.35, [[("연령", {"bold": True})]], size=14, font=F_MED)
image(s, X0 + 6.75, BODY_Y + 0.5, W - 7.0, 2.1, "s09_datalab_age.png")
tb(s, X0 + 6.75, BODY_Y + 2.7, W - 7.0, 0.35, [[("1위 40대 · 2위 50대 · 3위 30대", {"bold": True, "color": LILAC})]], size=14)
b = box(s, X0, 5.35, W, 1.35, fill=LILAC_T)
shape_text(b, [[("인사이트  ", {"bold": True, "color": LILAC}), ("메타 타겟을 20대 중심이 아니라 30~50대 여성으로 잡아야 합니다", {})],
               [("인사이트  ", {"bold": True, "color": LILAC}), ("'나 어릴 때'라는 브랜드명의 향수 코드는 이 연령대에 특히 잘 맞습니다", {})]],
           size=15, margin=0.3, space_after=4)
notes(s, "카테고리 전체 데이터라는 점을 같이 말씀드립니다.")

# S10 시장 진단 ② 검색어와 시즌
L = (ASSETS / "s10_datalab_keywords.csv").read_text(encoding="utf-8-sig").splitlines()[8:]
rows = list(csv.DictReader(L))
tot = {k: sum(int(r[k]) for r in rows) for k in rows[0] if k != "날짜"}
mon = Counter()
for r in rows:
    mon[r["날짜"][:7]] += int(r["나어릴때아이라이너"])
kw_bar = [("워터프루프 아이라이너", "워터프루프아이라이너"), ("펜 아이라이너", "펜아이라이너"), ("나어릴때 아이라이너", "나어릴때아이라이너"),
          ("브러쉬 아이라이너", "브러쉬아이라이너"), ("엣지홀릭", "엣지홀릭")]
assert [tot[k] for _, k in kw_bar] == [7974, 2159, 1076, 454, 14], "CSV 검산 불일치"
months = ["2025-11", "2025-12", "2026-01", "2026-02", "2026-03", "2026-04", "2026-05", "2026-06", "2026-07", "2026-08", "2026-09"]
assert [mon[m] for m in months] == [212, 213, 59, 80, 63, 29, 115, 44, 40, 49, 83], "CSV 월별 검산 불일치"

s = content("02 사전 진단", "고객은 '워터프루프'로 찾고, 나어릴때는 작년 11~12월에 가장 많이 찾았습니다",
            "네이버 데이터랩 검색어 통계 · 아이라이너 카테고리 · 1년 상대지수",
            source="출처: 네이버 데이터랩. 수치는 실제 검색 횟수가 아닌 상대지수입니다.")
box(s, X0, BODY_Y, 5.9, 3.35, fill=WHITE, line=LINE)
cd = CategoryChartData()
cd.categories = [n for n, _ in kw_bar]
cd.add_series("1년 합계 상대지수", [tot[k] for _, k in kw_bar])
ch = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(X0 + 0.1), Inches(BODY_Y + 0.05), Inches(5.7), Inches(3.25), cd).chart
chart_font(ch, 11)
ch.has_title = True
fill_tf(ch.chart_title.text_frame, [[("1년 합계 상대지수", {"bold": True})]], size=13, font=F_MED)
ch.has_legend = False
ch.category_axis.reverse_order = True
ch.category_axis.format.line.fill.background()
ch.value_axis.visible = False
ch.value_axis.has_major_gridlines = False
ch.value_axis.maximum_scale = 9500
pl = ch.plots[0]
pl.gap_width = 60
pl.has_data_labels = True
pl.data_labels.number_format = "#,##0"
pl.data_labels.number_format_is_linked = False
pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
pl.data_labels.font.size = Pt(12)
pl.data_labels.font.bold = True
for i, col in enumerate([LILAC, LILAC_M, PEACH, LILAC_M, LILAC_M]):
    pt = pl.series[0].points[i]
    pt.format.fill.solid()
    pt.format.fill.fore_color.rgb = rgb(col)

box(s, X0 + 6.15, BODY_Y, W - 6.15, 3.35, fill=WHITE, line=LINE)
cd = CategoryChartData()
cd.categories = [f"{m[2:4]}.{m[5:]}" for m in months]
cd.add_series("나어릴때 아이라이너", [mon[m] for m in months])
ch = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(X0 + 6.25), Inches(BODY_Y + 0.05), Inches(W - 6.35), Inches(3.25), cd).chart
chart_font(ch, 10)
ch.has_title = True
fill_tf(ch.chart_title.text_frame, [[("월별 '나어릴때 아이라이너' 상대지수", {"bold": True})]], size=13, font=F_MED)
ch.has_legend = False
ch.value_axis.visible = False
ch.value_axis.has_major_gridlines = False
ch.value_axis.minimum_scale = 0
ch.value_axis.maximum_scale = 260
ch.category_axis.format.line.color.rgb = rgb(LINE)
ser = ch.plots[0].series[0]
ser.smooth = False
ser.format.line.color.rgb = rgb(LILAC)
ser.format.line.width = Pt(2.25)
ser.marker.format.fill.solid()
ser.marker.format.fill.fore_color.rgb = rgb(LILAC)
ser.marker.format.line.color.rgb = rgb(LILAC)
pl = ch.plots[0]
pl.has_data_labels = True
pl.data_labels.number_format = "0"
pl.data_labels.number_format_is_linked = False
pl.data_labels.position = XL_LABEL_POSITION.ABOVE
pl.data_labels.font.size = Pt(10)
for i in (0, 1):
    pt = ser.points[i]
    pt.data_label.font.bold = True
    pt.data_label.font.size = Pt(12)
    pt.data_label.font.color.rgb = rgb(LILAC)
b = box(s, X0, 5.45, W, 1.3, fill=LILAC_T)
shape_text(b, [
    [("작년 11~12월 브랜드 검색이 평소의 2~4배 → ", {}), ("이번 11월 슈퍼적립이 승부처입니다", {"bold": True, "color": LILAC})],
    [("가장 큰 검색어는 '워터프루프 아이라이너' → ", {}), ("All New 워터프루프 라인으로 공략합니다", {"bold": True, "color": LILAC})],
    [("제품명 '엣지홀릭'은 거의 검색되지 않습니다 → ", {}), ("고객은 브랜드명으로 찾습니다", {"bold": True, "color": LILAC})],
], size=14, margin=0.3, space_after=2)

# S11 네이버 진단 ① 브랜드 검색 결과
s = content("02 사전 진단", "'나어릴때'를 검색한 고객이 다른 곳으로 새고 있습니다", "네이버 통합검색 '나어릴때' · 2026.10.23")
image(s, X0, BODY_Y, 2.9, 3.85, "s11_search_mobile.png")
tb(s, X0, BODY_Y + 3.88, 2.9, 0.25, [[("모바일", {})]], size=11, color=CAP, align=PP_ALIGN.CENTER)
image(s, X0 + 3.1, BODY_Y, 4.3, 3.85, "s11_search_pc.png")
tb(s, X0 + 3.1, BODY_Y + 3.88, 4.3, 0.25, [[("PC", {})]], size=11, color=CAP, align=PP_ALIGN.CENTER)
calls = [
    ("브랜드검색 미집행", "PC·모바일 모두 브랜드 전용 영역이 없습니다"),
    ("AI 브리핑의 해석", "모바일 최상단 AI 브리핑이 브랜드명을 '어린 시절 추억'으로도 해석해, 90년대 추억 콘텐츠가 섞여 나옵니다"),
    ("다른 가격 노출", "AI 브리핑 관련 상품에 공식가(8,500원)와 다른 20,700원 판매처 상품이 노출됩니다"),
]
rx = X0 + 7.65
for i, (t, d) in enumerate(calls):
    y = BODY_Y + i * 1.3
    badge(s, rx, y, str(i + 1))
    tb(s, rx + 0.55, y, W - 8.2, 0.42, [[(t, {"bold": True})]], size=15, font=F_MED, anchor=MSO_ANCHOR.MIDDLE)
    tb(s, rx + 0.55, y + 0.45, W - 8.2, 0.8, [[(d, {})]], size=13, color=SUB)
b = box(s, X0, 6.1, W, 0.68, fill=LILAC_T)
shape_text(b, [[("브랜드명이 일반 단어라서, 브랜드검색으로 상단을 지키는 것이 특히 중요합니다. ", {"bold": True, "color": LILAC}),
                ("AI 검색 노출은 NICECORE SEO/AEO 진단으로 함께 점검할 수 있습니다.", {})]], size=13, margin=0.25)
notes(s, "검색 비중의 91%가 모바일이라 모바일 화면을 크게 보여줍니다.")

# S12 네이버 진단 ② 쇼핑 탭 광고 공백
s = content("02 사전 진단", "네이버는 나어릴때를 첫 번째로 보여주는데, 광고 자리는 비어 있습니다", "네이버 쇼핑 모바일 · 2026.10.23")
image(s, X0, BODY_Y, 2.4, 4.6, "s12_shop_waterproof.png")
tb(s, X0, BODY_Y + 4.62, 2.4, 0.22, [[("워터프루프 아이라이너", {})]], size=10, color=CAP, align=PP_ALIGN.CENTER)
image(s, X0 + 2.6, BODY_Y, 2.4, 4.6, "s12_shop_pen.png")
tb(s, X0 + 2.6, BODY_Y + 4.62, 2.4, 0.22, [[("펜 아이라이너", {})]], size=10, color=CAP, align=PP_ALIGN.CENTER)
rx = X0 + 5.3
hl = {5: LILAC_T}
table(s, rx, BODY_Y, [2.15, 2.49, 2.49], [
    ["영역", "워터프루프 아이라이너", "펜 아이라이너"],
    ["브랜드 필터", "나어릴때 첫 칸", "나어릴때 첫 칸"],
    ["상단 쇼핑 브랜드형 광고", "네오젠, 케이트", "케이트"],
    ["쇼핑검색 광고", "삐아 2, 머지 2", "머지 3, 비오엠, 패션캣"],
    ["하단 파워링크", "케이트, 에네스티 외", "블랑디바, 머지 외"],
    [[("나어릴때 광고", {"bold": True, "color": LILAC})], [("없음", {"bold": True, "color": LILAC})], [("없음", {"bold": True, "color": LILAC})]],
], row_h=0.48, size=13, fills=hl, bold_cols=(0,))
tb(s, rx, BODY_Y + 3.15, W - 5.3, 1.3, [
    [("두 키워드 모두 머지가 쇼핑검색과 파워링크를 함께 잡고 있습니다 (가격대 6,900~8,900원, 나어릴때와 비슷)", {})],
    [("네이버에서의 실제 경쟁사는 머지입니다", {"bold": True, "color": LILAC})],
], size=14, bullet=True, space_after=4)

# S13 네이버 진단 ③ 카탈로그와 가격
s = content("02 사전 진단", "오가닉 1위 자리에 광고만 붙이면, 두 번 노출됩니다", "'워터프루프 아이라이너' 쇼핑 탭 · 엣지홀릭 브러쉬 브랜드 카탈로그")
bigs = [
    ("4,901", "카탈로그 리뷰 수 (평점 4.57)", "광고 4개 바로 아래 오가닉 1번째. 여기에 쇼핑검색 광고를 더하면 광고·오가닉 이중 노출이 됩니다"),
    ("38", "판매처 수", "최저가 1개 7,060원 / 2개 15,200원 (공식가 8,500원, 1+1 17,000원). 광고로 데려온 고객이 다른 판매처에서 구매할 수 있습니다"),
]
for i, (n, l, d) in enumerate(bigs):
    y = BODY_Y + i * 1.9
    box(s, X0, y, 7.2, 1.75, fill=WHITE, line=LINE)
    tb(s, X0 + 0.3, y + 0.15, 2.3, 0.9, [[(n, {})]], size=48, font=F_BLACK, color=LILAC, anchor=MSO_ANCHOR.MIDDLE)
    tb(s, X0 + 0.3, y + 1.08, 2.3, 0.5, [[(l, {})]], size=12, color=SUB)
    tb(s, X0 + 2.75, y + 0.2, 4.2, 1.35, [[(d, {})]], size=14, anchor=MSO_ANCHOR.MIDDLE)
image(s, X0 + 7.5, BODY_Y, W - 7.5, 3.65, "s13_catalog.png")
b = box(s, X0, 5.85, W, 0.85, fill=LILAC_T)
shape_text(b, [[("쇼핑검색 광고는 가장 적은 비용으로 효과를 볼 수 있는 지점입니다. ", {"bold": True, "color": LILAC}),
                ("판매처 가격 정책은 함께 점검을 권해드립니다.", {})]], size=14, margin=0.3)
notes(s, "가격 이슈는 단정하지 않고 '점검 권장' 톤으로 말합니다.")

# S14 메타 진단 ① 나어릴때 vs 경쟁사
s = content("02 사전 진단", "경쟁사는 '제품력'을 보여주고, 나어릴때는 '할인'만 보여줍니다", "메타 광고 라이브러리 · 집행 중 광고 · 2026.10.23")
lw = 4.1
box(s, X0, BODY_Y, lw, 3.95, fill=PEACH_T)
tb(s, X0 + 0.2, BODY_Y + 0.1, lw - 0.4, 0.35, [[("나어릴때", {"bold": True})]], size=15, font=F_MED)
image(s, X0 + 0.2, BODY_Y + 0.5, 1.6, 1.9, "s14_meta_naeorilttae.png")
tb(s, X0 + 1.95, BODY_Y + 0.5, lw - 2.1, 3.35, [
    [("집행 중 광고 5개, 모두 10/11 시작", {})],
    [("할로윈 펜 아이라이너 할인(10/12~31) 단일 프로모션", {})],
    [("전부 일러스트 할인 배너, 제품 시연·영상 없음", {})],
    [("5개 중 4개가 같은 문안 → 소재별 비교 테스트 구조가 없습니다", {"bold": True})],
], size=12, bullet=True, space_after=3)
rx = X0 + lw + 0.25
rw = W - lw - 0.25
box(s, rx, BODY_Y, rw, 3.95, fill=WHITE, line=LINE)
tb(s, rx + 0.2, BODY_Y + 0.1, rw - 0.4, 0.35, [[("경쟁사 ", {"bold": True}), ("('아이라이너' 키워드 광고 약 110개, 절반 이상이 영상)", {"size": 12, "color": SUB})]], size=15, font=F_MED)
comp = [("s14_meta_clio.png", "클리오", "\"올영 1위 라이너\", 1+1 · 개당 6천원대"),
        ("s14_meta_sisterand.png", "씨스터앤", "\"비 오고 물 튀겨도\" 번짐 방지 영상, 6/23 시작 소재를 4개월째 집행"),
        ("s14_meta_naming.png", "네이밍", "\"속쌍도 번짐 걱정 없는 트임라이너, 1mm\"")]
cw = (rw - 0.4 - 0.4) / 3
for i, (f, n, d) in enumerate(comp):
    x = rx + 0.2 + i * (cw + 0.2)
    image(s, x, BODY_Y + 0.5, cw, 1.9, f)
    tb(s, x, BODY_Y + 2.5, cw, 1.35, [[(n, {"bold": True, "font": F_MED, "size": 13})], [(d, {})]], size=12, color=SUB, space_after=2)
tb(s, X0, 6.05, 4.1, 0.6, [[("경쟁 광고의 메시지는 다섯 가지로 모입니다", {"bold": True})]], size=14, font=F_MED, anchor=MSO_ANCHOR.MIDDLE)
msgs = ["두께 수치", "번짐 방지 시연", "용도(트임·애교살)", "1위·대란템", "1+1 가격"]
cx = X0 + 4.2
cws = [1.25, 1.55, 1.85, 1.35, 1.2]
for m_, w_ in zip(msgs, cws):
    chip(s, cx, 6.17, w_, m_, size=12, h=0.36)
    cx += w_ + 0.15
notes(s, "경쟁사 비판이 아니라 '시장에서 통하는 메시지'를 보여주는 장표입니다.")

# S15 메타 진단 ② 리뷰 속 강점
s = content("02 사전 진단", "경쟁사가 광고비로 말하는 강점을, 나어릴때 고객은 이미 말하고 있습니다", "나어릴때 자사몰 실제 고객 리뷰")
quotes = [("\"선명하게 잘 그려지고 번지지 않아서 더울 때도 열심히 사용했어요!\"", "오리지널 엣지홀릭 브러쉬 아이라이너"),
          ("\"번지지 않는!!\"", "엣지홀릭 펜 아이라이너"),
          ("\"아이라이너는 역시 나어릴때\"", "엣지홀릭 펜 아이라이너 (재구매 고객)")]
cw = (W - 0.6) / 3
for i, (q, src) in enumerate(quotes):
    x = X0 + i * (cw + 0.3)
    box(s, x, BODY_Y, cw, 2.05, fill=WHITE, line=LINE)
    tb(s, x + 0.25, BODY_Y + 0.05, 0.6, 0.6, [[("“", {})]], size=40, font=F_BLACK, color=LILAC_M)
    tb(s, x + 0.25, BODY_Y + 0.55, cw - 0.5, 0.95, [[(q, {"bold": True})]], size=15, font=F_MED, anchor=MSO_ANCHOR.MIDDLE)
    tb(s, x + 0.25, BODY_Y + 1.6, cw - 0.5, 0.35, [[("— " + src, {})]], size=12, color=SUB)
image(s, X0, BODY_Y + 2.3, 4.6, 2.5, "s15_reviews.png")
b = box(s, X0 + 4.85, BODY_Y + 2.3, W - 4.85, 2.5, fill=LILAC)
shape_text(b, [[("'번짐 방지', '선명함', '재구매'는 경쟁사가 가장 많이 쓰는 메시지입니다.", {"bold": True})],
               [("나어릴때는 이걸 실제 리뷰로 갖고 있지만, 지금 광고에는 쓰이지 않고 있습니다.", {})]],
           size=17, color=WHITE, margin=0.4, space_after=6)
notes(s, "4번 진단 중 가장 강한 장표. 기존 대행사의 '분석·역제안 부족'을 직접 말하지 않고 보여줍니다.")

# S16 인스타그램·자사몰 진단
s = content("02 사전 진단", "광고를 보고 들어온 고객이, 멈춘 계정을 보게 됩니다")
lw = 7.15
box(s, X0, BODY_Y - 0.15, lw, 4.95, fill=WHITE, line=LINE)
tb(s, X0 + 0.25, BODY_Y - 0.03, lw, 0.38, [[("인스타그램 @naeorilttae_official", {"bold": True})]], size=15, font=F_MED)
image(s, X0 + 0.25, BODY_Y + 0.45, 2.4, 4.2, "s16_insta_feed.png")
tb(s, X0 + 2.85, BODY_Y + 0.45, lw - 3.05, 4.2, [
    [("팔로워 1,615명 · 게시물 79개", {})],
    [("마지막 게시물 9/12 → ", {"bold": True}), ("40일 이상 업로드 없음", {"bold": True, "color": LILAC})],
    [("8/9 하루 7개, 8/22 하루 3개 → 꾸준히 올리기보다 몰아서 올리는 패턴", {})],
    [("10/11부터 할로윈 광고가 도는데, 피드에는 관련 게시물이 없습니다", {})],
    [("영문 구계정(@naeorilttae, 팔로워 95명)이 따로 남아 있어 검색 시 혼선이 있을 수 있습니다", {})],
], size=13, bullet=True, space_after=5)
rx = X0 + lw + 0.25
rw = W - lw - 0.25
box(s, rx, BODY_Y - 0.15, rw, 4.95, fill=WHITE, line=LINE)
tb(s, rx + 0.25, BODY_Y - 0.03, rw, 0.38, [[("자사몰 광고 추적 설정", {"bold": True})]], size=15, font=F_MED)
for i, (t, ok) in enumerate([("메타 픽셀", True), ("네이버 프리미엄로그분석", True), ("GA4", False), ("카카오 픽셀", False)]):
    y = BODY_Y + 0.5 + i * 0.62
    tb(s, rx + 0.25, y, rw - 1.8, 0.45, [[(t, {})]], size=14, anchor=MSO_ANCHOR.MIDDLE)
    chip(s, rx + rw - 1.45, y + 0.06, 1.2, "설치됨" if ok else "없음", fill=LILAC if ok else PEACH, color=WHITE if ok else TEXT, h=0.34)
    hline(s, rx + 0.25, y + 0.55, rw - 0.5)
b = box(s, rx + 0.25, BODY_Y + 3.15, rw - 0.5, 1.45, fill=LILAC_T)
shape_text(b, [[("채널별 성과를 같은 기준으로 비교하려면 GA4 설치를 권해드립니다", {"bold": True, "color": LILAC})]], size=14, margin=0.2)

# S17 진단 종합
s = content("02 사전 진단", "새로 만들 필요 없이, 이미 있는 것을 연결하면 됩니다")
sums = [("비어 있는 검색 자리", "네이버가 첫 번째로 보여주는 브랜드인데, 광고 자리는 머지·삐아가 차지", "워터프루프 쇼핑검색 + 브랜드검색으로 되찾기"),
        ("쓰이지 않는 리뷰", "'안 번진다'는 리뷰가 있는데, 광고는 할인 배너뿐", "리뷰 기반 소재로 메타 재설계"),
        ("끊어진 연결", "광고는 돌고, 계정은 멈추고, 추적은 일부만", "인스타 정기 운영 + GA4로 광고와 계정, 측정을 하나로")]
cw = (W - 0.6) / 3
for i, (t, p, sol) in enumerate(sums):
    x = X0 + i * (cw + 0.3)
    box(s, x, BODY_Y - 0.15, cw, 4.85, fill=WHITE, line=LINE)
    tb(s, x + 0.3, BODY_Y + 0.05, 2, 0.9, [[(f"0{i + 1}", {})]], size=44, font=F_BLACK, color=LILAC)
    tb(s, x + 0.3, BODY_Y + 1.05, cw - 0.6, 0.45, [[(t, {"bold": True})]], size=19, font=F_MED)
    tb(s, x + 0.3, BODY_Y + 1.65, cw - 0.6, 1.2, [[(p, {})]], size=14, color=SUB)
    b = box(s, x + 0.2, BODY_Y + 3.15, cw - 0.4, 1.35, fill=LILAC_T)
    shape_text(b, [[("→ " + sol, {"bold": True, "color": LILAC})]], size=14, margin=0.2)

# S18 간지
section_slide("03", "운영 전략", "네이버는 지키고, 메타는 키웁니다",
              ["채널 전략 · 네이버 실행안", "메타 캠페인 구조 · 소재 방향", "11월 슈퍼적립 집중 플랜"])

# S19 채널 전략
s = content("03 운영 전략", "채널마다 역할을 나누고, 각자의 지표로 관리합니다")
table(s, X0, BODY_Y, [2.6, 4.4, 2.5, 2.93], [
    ["채널", "역할", "랜딩", "핵심 지표"],
    ["네이버 검색광고", "브랜드·핵심 키워드 자리 확보", "브랜드스토어", "ROAS, 노출 점유"],
    ["네이버 GFA", "관심 고객 재방문, 슈퍼적립 노출", "브랜드스토어", "ROAS"],
    ["메타", "30~50대 신규 고객 확보", "자사몰", "구매 ROAS, CPA"],
    ["인스타그램", "광고 유입 고객의 신뢰 확보", "계정", "게시 빈도, 반응"],
], row_h=0.66, size=15, bold_cols=(0,))
b = box(s, X0, 5.6, W, 0.95, fill=LILAC_T)
shape_text(b, [[("네이버는 '이미 찾는 고객'을 놓치지 않는 역할, ", {"bold": True, "color": LILAC}),
                ("메타는 '아직 모르는 고객'을 데려오는 역할입니다", {"bold": True, "color": LILAC})]], size=16, margin=0.3)

# S20 네이버 실행안
s = content("03 운영 전략", "가장 효과가 큰 순서대로 실행합니다")
pri = [("1순위", "워터프루프 아이라이너 쇼핑검색", ["오가닉 1위 카탈로그에 광고를 더해 이중 노출", "머지·삐아가 차지한 광고 자리 진입"]),
       ("2순위", "브랜드검색 + 브랜드 키워드 파워링크", ["'나어릴때' 검색 상단 확보, 추억 콘텐츠·타 판매처로의 이탈 방어", "작년 11~12월 브랜드 검색 피크 시기에 맞춰 집행"]),
       ("3순위", "용도 키워드 확장", ["펜 아이라이너, 브러쉬 아이라이너, 리퀴드 아이라이너, 애교살 라이너", "데이터랩 인기검색어 기반으로 매월 추가·정리"])]
lw = 8.1
for i, (tag, t, items) in enumerate(pri):
    y = BODY_Y - 0.15 + i * 1.65
    box(s, X0, y, lw, 1.5, fill=WHITE, line=LINE)
    chip(s, X0 + 0.25, y + 0.2, 0.95, tag, fill=LILAC if i == 0 else LILAC_M, color=WHITE, h=0.34)
    tb(s, X0 + 1.35, y + 0.17, lw - 1.6, 0.4, [[(t, {"bold": True})]], size=17, font=F_MED, anchor=MSO_ANCHOR.MIDDLE)
    tb(s, X0 + 1.35, y + 0.65, lw - 1.6, 0.8, [[(it, {})] for it in items], size=13, color=SUB, bullet=True, space_after=2)
rx = X0 + lw + 0.3
rw = W - lw - 0.3
b = box(s, rx, BODY_Y - 0.15, rw, 4.8, fill=LILAC_T)
tb(s, rx + 0.3, BODY_Y + 0.1, rw - 0.6, 0.45, [[("GFA", {})]], size=22, font=F_BLACK, color=LILAC)
tb(s, rx + 0.3, BODY_Y + 0.75, rw - 0.6, 3.5, [[("브랜드스토어 상세 조회·장바구니 고객 리타겟팅", {})], [("슈퍼적립 기간 노출 확대", {})], [("소재 2~3종 비교 테스트", {})]],
   size=15, bullet=True, space_after=10)
notes(s, "네이버는 대행수수료가 없다는 점을 S31에서 같이 말씀드립니다.")

# S21 메타 실행안 ① 캠페인 구조
s = content("03 운영 전략", "새 고객과 돌아올 고객을 나눠서 운영합니다")
root = box(s, X0, BODY_Y + 0.4, 1.7, 3.05, fill=LILAC)
shape_text(root, [[("메타", {"bold": True})], [("캠페인", {})]], size=18, color=WHITE, align=PP_ALIGN.CENTER)
camps = [("캠페인 A — 신규 고객", [("타겟", "30~50대 여성, 메이크업·뷰티 관심사, 구매 고객 유사 타겟"), ("소재", "제품력 중심 (번짐 테스트, 라인 시연)"), ("지표", "CPA, 구매 ROAS")], LILAC),
         ("캠페인 B — 리타겟팅", [("타겟", "자사몰 방문 30일, 장바구니 7일"), ("소재", "혜택 중심 (1+1, 기간 한정)"), ("지표", "구매 ROAS")], LILAC_M)]
bx = X0 + 2.3
bw = W - 2.3
for i, (t, rows_, col) in enumerate(camps):
    y = BODY_Y - 0.15 + i * 2.05
    c = s.shapes.add_connector(MSO_CONNECTOR.ELBOW, Inches(X0 + 1.7), Inches(BODY_Y + 1.92), Inches(bx), Inches(y + 0.9))
    _strip_style(c)
    c.line.color.rgb = rgb(LILAC_M)
    c.line.width = Pt(1.5)
    box(s, bx, y, bw, 1.85, fill=WHITE, line=col, line_w=1.25)
    tb(s, bx + 0.25, y + 0.12, bw - 0.5, 0.4, [[(t, {"bold": True, "color": LILAC})]], size=16, font=F_MED)
    cw = (bw - 0.5 - 0.4) / 3
    for j, (k, v) in enumerate(rows_):
        x = bx + 0.25 + j * (cw + 0.2)
        chip(s, x, y + 0.65, 0.7, k, h=0.3, size=11)
        tb(s, x, y + 1.02, cw, 0.8, [[(v, {})]], size=13)
b = box(s, X0, 5.95, W, 0.8, fill=PEACH_T)
shape_text(b, [[("측정 기반  ", {"bold": True, "color": LILAC}), ("메타 픽셀 구매 이벤트 점검 및 전환 API 연동 검토  ·  GA4 설치로 채널 간 성과 비교", {})]], size=14, margin=0.3)

# S22 메타 실행안 ② 소재 방향
s = content("03 운영 전략", "고객 리뷰에서 소재를 뽑습니다")
cons = [("번짐 테스트", "15초 영상", "\"더울 때도, 하루 종일. 번지지 않는 라인\"", "리뷰 \"번지지 않아서 더울 때도 열심히 사용했어요\""),
        ("라인 시연", "릴스 · 클로즈업 영상", "\"꼬리 라인, 브러쉬 한 번이면 끝\"", "경쟁 광고의 절반 이상이 영상"),
        ("브러쉬 vs 펜", "비교 이미지 · 카루셀", "\"붓처럼 vs 펜처럼, 나한테 맞는 라인은?\"", "두 제품 모두 리뷰 700건 이상"),
        ("브랜드 향수", "이미지 · 영상", "\"나 어릴 때 그 눈매, 지금 다시\"", "주 고객층 40~50대")]
cw = (W - 0.9) / 4
for i, (t, fmt, copy_, why) in enumerate(cons):
    x = X0 + i * (cw + 0.3)
    box(s, x, BODY_Y - 0.15, cw, 4.15, fill=WHITE, line=LINE)
    badge(s, x + 0.25, BODY_Y + 0.05, str(i + 1))
    tb(s, x + 0.25, BODY_Y + 0.6, cw - 0.5, 0.42, [[(t, {"bold": True})]], size=18, font=F_MED)
    chip(s, x + 0.25, BODY_Y + 1.1, cw - 0.5, fmt, h=0.32, size=12)
    tb(s, x + 0.25, BODY_Y + 1.6, cw - 0.5, 1.25, [[(copy_, {"bold": True, "color": LILAC})]], size=15, font=F_MED, anchor=MSO_ANCHOR.MIDDLE)
    hline(s, x + 0.25, BODY_Y + 2.95, cw - 0.5)
    tb(s, x + 0.25, BODY_Y + 3.05, cw - 0.5, 0.85, [[("근거  ", {"bold": True}), (why, {})]], size=12, color=SUB)
b = box(s, X0, 6.15, W, 0.6, fill=PEACH_T)
shape_text(b, [[("혜택 소재(1+1, 슈퍼적립)는 리타겟팅 캠페인에서 별도 운영합니다", {})]], size=14, margin=0.3)
notes(s, "카피는 방향 예시이며, 실제 집행 문안은 협의 후 확정합니다.")

# S23 11월 슈퍼적립 집중 플랜
s = content("03 운영 전략", "11월 슈퍼적립부터 바로 효과를 내겠습니다", "〔확인〕슈퍼적립 기간·참여 상품 확정 후 세부 일정 반영")
steps = [("1주차 (11/2~)", "세팅", ["브랜드검색 신청, 워터프루프 쇼핑검색 집행", "메타 캠페인 구조 재편, 픽셀 점검", "신규 소재 1차 4종 제작"]),
         ("2~3주차", "집중", ["슈퍼적립 기간 GFA 노출 확대, 쇼핑검색 입찰 조정", "메타 리타겟팅에 슈퍼적립 혜택 소재 투입", "인스타 피드에 행사 콘텐츠 연동"]),
         ("4주차", "확장", ["1차 테스트 결과로 위닝 소재 선정 → 예산 이동", "반응 좋은 키워드 확장"]),
         ("12월 초", "리뷰", ["11월 성과 리포트 + 12월 연말 시즌 제안"])]
cw = (W - 0.9) / 4
hline(s, X0 + 0.2, BODY_Y + 0.25, W - 0.4, color=LILAC_M, width=2)
for i, (when, t, items) in enumerate(steps):
    x = X0 + i * (cw + 0.3)
    box(s, x + 0.1, BODY_Y + 0.1, 0.3, 0.3, fill=LILAC if i < 3 else LILAC_M, shape=MSO_SHAPE.OVAL)
    tb(s, x, BODY_Y + 0.55, cw, 0.4, [[(when, {"bold": True, "color": LILAC})]], size=16, font=F_MED)
    box(s, x, BODY_Y + 1.05, cw, 2.75, fill=WHITE, line=LINE)
    tb(s, x + 0.25, BODY_Y + 1.2, cw - 0.5, 0.45, [[(t, {"bold": True})]], size=20, font=F_BLACK)
    tb(s, x + 0.25, BODY_Y + 1.8, cw - 0.5, 1.95, [[(it, {})] for it in items], size=13, bullet=True, space_after=6)
CHECKS.append("S23 부제: 슈퍼적립 기간·참여 상품 확정 후 세부 일정 반영")

# S24 간지
section_slide("04", "운영 방식", "말씀하신 다섯 가지에 하나씩 답합니다",
              ["역제안 체계 · 위닝 소재 운영 · 소재 제작", "인스타그램 운영 · 리포트 샘플 · 소통 체계"])

# S25 역제안 체계
s = content("04 운영 방식", "요청하시기 전에, 먼저 가져갑니다", "답변 03 — 역제안")
proc = [("모니터링", "매주 데이터랩 인기검색어·급등 키워드 확인"), ("후보 선정", "나어릴때 제품과 맞는 키워드·소구점 선별"),
        ("제안", "주간 리포트에 '이번 주 제안' 항목으로 전달"), ("반영", "컨펌 후 키워드·소재에 바로 적용")]
cw = (W - 3 * 0.45) / 4
for i, (t, d) in enumerate(proc):
    x = X0 + i * (cw + 0.45)
    box(s, x, BODY_Y - 0.1, cw, 1.4, fill=WHITE, line=LINE)
    tb(s, x + 0.2, BODY_Y, cw - 0.4, 0.4, [[(f"0{i + 1} ", {"color": LILAC, "font": F_BLACK}), (t, {"bold": True})]], size=16, font=F_MED)
    tb(s, x + 0.2, BODY_Y + 0.45, cw - 0.4, 0.8, [[(d, {})]], size=13, color=SUB)
    if i < 3:
        arrow(s, x + cw + 0.08, BODY_Y + 0.45, 0.3, 0.3)
b = box(s, X0, BODY_Y + 1.45, W, 0.5, fill=PEACH_T)
shape_text(b, [[("예시  ", {"bold": True, "color": LILAC}), ("현재 인기검색어 '애교살 라이너' → 브러쉬 아이라이너 애교살 연출 소재 제안", {})]], size=14, margin=0.25)
tb(s, X0, BODY_Y + 2.1, 4, 0.35, [[("시즌 캘린더", {"bold": True})]], size=15, font=F_MED)
cal = [("11월", "슈퍼적립, 블랙프라이데이"), ("12월", "연말 선물, 1+1 세트"), ("1~3월", "새해·졸업·입학, 신학기 메이크업"),
       ("4~5월", "워터프루프 검색 피크 (데이터랩 5월 최고)"), ("6~8월", "여름 번짐 방지"), ("9~10월", "가을 브라운 컬러, 할로윈")]
cw = (W - 5 * 0.15) / 6
for i, (m_, d) in enumerate(cal):
    x = X0 + i * (cw + 0.15)
    chip(s, x, BODY_Y + 2.55, cw, m_, fill=LILAC if i < 2 else LILAC_M, color=WHITE, h=0.38, size=13)
    box(s, x, BODY_Y + 3.0, cw, 1.2, fill=WHITE, line=LINE)
    tb(s, x + 0.15, BODY_Y + 3.12, cw - 0.3, 1.0, [[(d, {})]], size=13)

# S26 위닝 소재 운영
s = content("04 운영 방식", "소재는 감으로 고르지 않고, 테스트로 고르고 키웁니다", "답변 02 — 분석")
ps = [("테스트", "컨셉이 다른 소재 4~6개를 같은 조건으로 동시 집행 (1주)"),
      ("판별", "CTR(관심) → CVR(구매 전환) → ROAS(매출 효율) 순서로 판단, 계정 평균 대비 비교"),
      ("확장", "위닝 소재에 예산 이동, 같은 컨셉의 변형 소재 추가 제작")]
cw = (W - 2 * 0.6) / 3
for i, (t, d) in enumerate(ps):
    x = X0 + i * (cw + 0.6)
    box(s, x, BODY_Y, cw, 3.0, fill=WHITE if i < 2 else LILAC_T, line=LINE if i < 2 else None)
    tb(s, x + 0.3, BODY_Y + 0.25, 2, 0.8, [[(f"0{i + 1}", {})]], size=36, font=F_BLACK, color=LILAC)
    tb(s, x + 0.3, BODY_Y + 1.05, cw - 0.6, 0.45, [[(t, {"bold": True})]], size=20, font=F_MED)
    tb(s, x + 0.3, BODY_Y + 1.6, cw - 0.6, 1.3, [[(d, {})]], size=14, color=SUB)
    if i < 2:
        arrow(s, x + cw + 0.13, BODY_Y + 1.3, 0.34, 0.4)
b = box(s, X0, 5.3, W, 1.0, fill=PEACH_T)
shape_text(b, [[("씨스터앤의 번짐 방지 영상은 6/23부터 4개월째 집행 중 — ", {}), ("오래 도는 소재가 검증된 소재입니다", {"bold": True, "color": LILAC})]], size=15, margin=0.3)

# S27 소재 제작
s = content("04 운영 방식", "소재 제작비는 따로 받지 않습니다", "답변 04 — 소재 제작 · 메타·GFA 소재 기준")
bigs = [("월 20건 내외", "기본 제작 물량"), ("최대 50건", "테스트·시즌 집중 시"), ("0원", "추가 제작비")]
cw = (W - 0.6) / 3
for i, (n, l) in enumerate(bigs):
    x = X0 + i * (cw + 0.3)
    box(s, x, BODY_Y - 0.1, cw, 1.6, fill=LILAC if i == 2 else WHITE, line=None if i == 2 else LINE)
    tb(s, x + 0.3, BODY_Y + 0.05, cw - 0.6, 0.85, [[(n, {})]], size=34, font=F_BLACK, color=WHITE if i == 2 else LILAC, anchor=MSO_ANCHOR.MIDDLE)
    tb(s, x + 0.3, BODY_Y + 0.95, cw - 0.6, 0.4, [[(l, {})]], size=14, color=WHITE if i == 2 else SUB)
tb(s, X0, BODY_Y + 1.8, 4.3, 2.0, [[("기존 제품 컷을 AI로 보정해 촬영 없이도 소재 수를 늘립니다", {"bold": True})], [("(보정 전후 예시)", {"color": SUB, "size": 13})]],
   size=17, font=F_MED, space_after=4)
tb(s, X0, 6.35, 4.3, 0.35, [[("소재 난이도에 따라 건수는 조정될 수 있습니다", {})]], size=12, color=CAP)
iw = (W - 4.6 - 0.3) / 2
for i, (f, l) in enumerate([("s27_ai_before.png", "보정 전"), ("s27_ai_after.png", "보정 후")]):
    x = X0 + 4.6 + i * (iw + 0.3)
    chip(s, x, BODY_Y + 1.8, 1.0, l, fill=LILAC if i else LILAC_M, color=WHITE, h=0.32)
    image(s, x, BODY_Y + 2.2, iw, 2.6, f)

# S28 인스타그램 운영
s = content("04 운영 방식", "광고와 같은 톤으로, 끊기지 않게 운영합니다", "답변 05 — 인스타그램")
lw = 5.8
blocks = [("운영 범위", ["주 2회 게시 (월 8~9건)", "기획 · 이미지 제작 · 문구 작성 · 업로드"]),
          ("광고 연동", ["광고 시즌과 같은 주제로 피드 구성 (광고를 보고 들어온 고객이 같은 메시지를 보도록)"])]
tb(s, X0, BODY_Y - 0.1, lw, 0.35, [[("운영 범위", {"bold": True, "color": LILAC})]], size=15, font=F_MED)
tb(s, X0, BODY_Y + 0.28, lw, 0.7, [[("주 2회 게시 (월 8~9건)", {})], [("기획 · 이미지 제작 · 문구 작성 · 업로드", {})]], size=14, bullet=True, space_after=2)
tb(s, X0, BODY_Y + 1.1, lw, 0.35, [[("진행 프로세스", {"bold": True, "color": LILAC})]], size=15, font=F_MED)
flow = ["월간 콘텐츠 계획 공유", "주간 시안 컨펌", "업로드", "월간 반응 리포트"]
fw = (lw - 3 * 0.25) / 4
for i, t in enumerate(flow):
    x = X0 + i * (fw + 0.25)
    b = box(s, x, BODY_Y + 1.5, fw, 0.75, fill=WHITE, line=LINE)
    shape_text(b, [[(t, {})]], size=12, align=PP_ALIGN.CENTER, margin=0.05)
    if i < 3:
        tb(s, x + fw, BODY_Y + 1.5, 0.25, 0.75, [[("›", {"bold": True})]], size=18, color=LILAC_M, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
tb(s, X0, BODY_Y + 2.45, lw, 0.35, [[("광고 연동", {"bold": True, "color": LILAC})]], size=15, font=F_MED)
tb(s, X0, BODY_Y + 2.83, lw, 0.75, [[("광고 시즌과 같은 주제로 피드 구성 (광고를 보고 들어온 고객이 같은 메시지를 보도록)", {})]], size=14)
b = box(s, X0, BODY_Y + 3.7, lw, 1.1, fill=LILAC)
shape_text(b, [[("비용  ", {"size": 15}), ("월 350만 원", {"bold": True, "size": 26, "font": F_BLACK})]], size=15, color=WHITE, margin=0.3)
rx = X0 + lw + 0.4
fw = (W - lw - 0.4 - 0.4) / 3
for i in range(3):
    image(s, rx + i * (fw + 0.2), BODY_Y + 0.5, fw, fw * 1.25, f"s28_feed_{i + 1}.png")
tb(s, rx, BODY_Y + 0.5 + fw * 1.25 + 0.1, W - lw - 0.4, 0.3, [[("피드 시안", {})]], size=12, color=CAP, align=PP_ALIGN.CENTER)
# 원고 note("협력사 언급 금지")는 작성 지침이라 발표자 노트에 넣지 않음

# S29 리포트 샘플
s = content("04 운영 방식", "숫자가 아니라 '왜'와 '다음'을 씁니다", "답변 02 — 분석 · 실제 운영 중인 주간 리포트 (광고주명 비공개)")
image(s, X0, BODY_Y - 0.1, 8.0, 4.85, "s29_report_sample.png")
rx = X0 + 8.3
for i, t in enumerate(["지난주 대비 변화와 원인", "매체·소재별 효율 비교", "이번 주 액션과 다음 주 제안"]):
    y = BODY_Y + i * 1.6
    box(s, rx, y, W - 8.3, 1.35, fill=WHITE if i < 2 else LILAC_T, line=LINE if i < 2 else None)
    badge(s, rx + 0.25, y + 0.47, str(i + 1))
    tb(s, rx + 0.85, y, W - 9.3, 1.35, [[(t, {"bold": True})]], size=16, font=F_MED, anchor=MSO_ANCHOR.MIDDLE)

# S30 소통 체계
s = content("04 운영 방식", "업무시간 안에는 바로 응대합니다", "답변 01 — 소통")
comm = [("전담 담당자", "김현수 팀장 + 〔확인〕실무 담당자 1명"), ("실시간 채널", "카카오톡 단톡방 (평일 〔확인〕09:00~18:00 응대)"),
        ("응대 기준", "〔확인〕업무시간 내 1시간 안에 1차 회신"), ("정기 소통", "주간 리포트 + 월간 리뷰 미팅")]
cw, chh = (W - 0.3) / 2, 2.1
for i, (t, d) in enumerate(comm):
    x, y = X0 + (i % 2) * (cw + 0.3), BODY_Y + (i // 2) * (chh + 0.3)
    box(s, x, y, cw, chh, fill=WHITE, line=LINE)
    badge(s, x + 0.35, y + 0.35, str(i + 1), d=0.5, size=15)
    tb(s, x + 1.1, y + 0.35, cw - 1.4, 0.5, [[(t, {"bold": True})]], size=20, font=F_MED, anchor=MSO_ANCHOR.MIDDLE)
    tb(s, x + 1.1, y + 1.0, cw - 1.4, 0.9, [[(d, {})]], size=16)
CHECKS += ["S30 실무 담당자 1명", "S30 응대 시간 09:00~18:00", "S30 업무시간 내 1시간 안에 1차 회신"]

# S31 비용
s = content("05 비용과 일정", "숨은 비용 없이, 한 장으로 정리했습니다")
tb(s, X0, BODY_Y - 0.15, 5, 0.35, [[("운영 조건", {"bold": True, "color": LILAC})]], size=15, font=F_MED)
table(s, X0, BODY_Y + 0.25, [2.5, 3.6], [
    ["항목", "조건"],
    ["네이버 검색광고 · GFA", [("대행수수료 없음", {"bold": True, "color": LILAC}), (" (네이버 공식대행)", {})]],
    ["메타", "광고 집행비의 15%"],
    ["소재 제작 (메타·GFA)", [("추가 비용 없음", {"bold": True, "color": LILAC}), (" (월 20건 내외, 최대 50건)", {})]],
    ["인스타그램 운영", "월 350만 원 (주 2회)"],
    ["계약 기간", "1년 단위 기본"],
], row_h=0.58, size=13, bold_cols=(0,))
rx = X0 + 6.45
tb(s, rx, BODY_Y - 0.15, 6, 0.35, [[("월 예산별 배분 예시 ", {"bold": True, "color": LILAC}), ("(단위 만 원)", {"size": 12, "color": SUB})]], size=15, font=F_MED)
R = {c: PP_ALIGN.RIGHT for c in range(1, 5)}
table(s, rx, BODY_Y + 0.25, [1.15, 1.25, 1.0, 1.0, 1.58], [
    ["월 광고비", "네이버 검색광고", "GFA", "메타", "메타 수수료(15%)"],
    ["300", "120", "60", "120", "18"],
    ["500", "170", "100", "230", "34.5"],
    ["1,000", "300", "200", "500", "75"],
], row_h=0.58, size=13, bold_cols=(0,), aligns=R)
tb(s, X0, 6.05, W, 0.7, [[("배분은 예시이며, 목표와 성과에 따라 조정합니다. 브랜드검색은 조회수 기준 별도 단가로 산정됩니다. 〔확인〕VAT 별도 여부", {})]], size=12, color=SUB)
CHECKS.append("S31 VAT 별도 여부")

# S32 일정과 마무리
s = content("05 비용과 일정", "11월 첫 주부터 시작할 수 있습니다")
tl = [("10/30", "제안"), ("11월 1주", "계정 연결·세팅 (브랜드검색, 워터프루프 쇼핑검색 우선)"), ("11월 2~4주", "슈퍼적립 집중 운영"), ("12월 초", "11월 성과 리뷰 + 연말 시즌 제안")]
cw = (W - 0.9) / 4
hline(s, X0 + 0.2, BODY_Y + 0.15, W - 0.4, color=LILAC_M, width=2)
for i, (w_, d) in enumerate(tl):
    x = X0 + i * (cw + 0.3)
    box(s, x + 0.05, BODY_Y, 0.3, 0.3, fill=LILAC, shape=MSO_SHAPE.OVAL)
    tb(s, x, BODY_Y + 0.45, cw, 0.4, [[(w_, {"bold": True, "color": LILAC})]], size=17, font=F_MED)
    tb(s, x, BODY_Y + 0.9, cw - 0.1, 0.9, [[(d, {})]], size=14)
b = box(s, X0, 4.2, W, 1.2, fill=LILAC)
shape_text(b, [[("이미 가진 강점을, 광고가 제대로 쓰게 만들겠습니다", {})]], size=26, font=F_BLACK, color=WHITE, align=PP_ALIGN.CENTER)
tb(s, X0, 5.75, 9, 0.45, [[("NICECORE 김현수 팀장 / 〔확인〕연락처 / 〔확인〕이메일", {})]], size=16, font=F_MED, anchor=MSO_ANCHOR.MIDDLE)
s.shapes.add_picture(str(ASSETS / "logo_nicecore.png"), Inches(X1 - 2.0), Inches(5.85), Inches(2.0), Inches(0.29))
CHECKS += ["S01·S32 담당자 연락처", "S01·S32 담당자 이메일"]

prs.save(str(OUT))
print("saved", OUT, "slides", len(prs.slides._sldIdLst))
print("missing images", len(MISSING), sorted(MISSING))
